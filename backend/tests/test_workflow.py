import io
import json
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr

from app.config import Settings
from app.database import Database, Job, Quote, uid
from app.main import create_app
from app.schemas import QuoteFields
from app.repository import record, set_fields
from app.upstage import AgentError, UpstageClient


def xlsx(quantity=10, amount=250000, suffix=""):
    wb = Workbook()
    ws = wb.active
    for row in [["견적번호", "T-001"+suffix], ["고객사", "가상제조사"], ["발행 회사", "가상공급사"], ["작성일", "2026-10-16"], ["통화", "KRW"], ["결제 조건", "납품 후 30일"], ["세율", 10], ["공급가액", 250000], ["부가세", 25000], ["총액", 275000], [], ["품목명", "규격", "단위", "수량", "단가", "금액"], ["온도 센서", "TS-100", "개", quantity, 25000, amount]]:
        ws.append(row)
    data = io.BytesIO()
    wb.save(data)
    wb.close()
    return data.getvalue()


@pytest.fixture
def app(tmp_path):
    return create_app(Settings(_env_file=None, data_dir=tmp_path, extraction_provider="local"))


@pytest.fixture
def client(app):
    with TestClient(app) as c:
        yield c


def upload(client, data=None, filename="sample.xlsx"):
    r = client.post("/api/add", files={"file": (filename, data if data is not None else xlsx())})
    assert r.status_code == 201, r.text
    return r.json()


def extract(client, quote_id):
    r = client.post("/api/extraction-jobs", json={"quote_id": quote_id, "provider": "local"})
    assert r.status_code == 202, r.text
    job = client.get("/api/jobs/"+r.json()["job_id"]).json()
    assert job["status"] == "completed", job
    return client.get("/api/quotes/"+quote_id).json()


def test_upload_extract_search_edit_history_and_restart(client, app):
    original_bytes = xlsx()
    q = upload(client, original_bytes)
    assert q["quote_id"] != "stub-quote-id"
    before = client.get("/api/extract", params={"quote_id": q["quote_id"]}).json()
    assert before["extraction_status"] == "pending"  # GET has no side effect.
    q = extract(client, q["quote_id"])
    assert q["customer_name"] == "가상제조사"
    assert q["total_amount"] == "275000"
    assert q["items"][0]["quantity"] == "10"
    extraction = client.get("/api/extract", params={"quote_id": q["quote_id"]}).json()
    assert extraction["source_text"] and extraction["latest_job"]["status"] == "completed"
    assert client.get(f"/api/quotes/{q['quote_id']}/jobs").json()["total"] == 1
    assert client.post("/api/search", json={"item_name": "센서", "amount_min": "275000", "date_from": "2026-10-01"}).json()["total"] == 1
    assert client.get("/api/search", params={"keyword": "가상제조사"}).json()["total"] == 1
    result = client.patch("/api/edit/"+q["quote_id"], json={"title": "수정", "expected_version": q["version"]}).json()
    assert result["items"][0]["quantity"] == "10"  # Omitted list preserved.
    assert client.patch("/api/edit/"+q["quote_id"], json={"title": "충돌", "expected_version": q["version"]}).status_code == 409
    cleared = client.patch("/api/edit/"+q["quote_id"], json={"title": None, "items": None}).json()
    assert cleared["title"] is None and cleared["items"] == []
    assert client.get(f"/api/quotes/{q['quote_id']}/history").json()["total"] >= 4
    # A second app/engine points at the same persisted DB and files.
    with TestClient(create_app(app.state.config)) as restarted:
        assert restarted.get("/api/quotes/"+q["quote_id"]).json()["title"] is None
        assert restarted.get(f"/api/quotes/{q['quote_id']}/source").content == original_bytes


def test_original_amount_and_missing_quantity_preserved(client):
    q = extract(client, upload(client, xlsx(quantity=None, amount=190000))["quote_id"])
    assert q["items"][0]["quantity"] is None
    assert q["items"][0]["amount"] == "190000"
    assert any(w["code"] == "subtotal_mismatch" for w in q["validation"]["warnings"])


def test_zero_is_not_missing(client):
    q = extract(client, upload(client, xlsx(quantity=0, amount=0))["quote_id"])
    assert q["items"][0]["quantity"] == "0" and q["items"][0]["amount"] == "0"


def test_duplicate_trash_restore_source(client):
    original_bytes = xlsx()
    q = upload(client, original_bytes)
    assert upload(client, original_bytes)["duplicate"] is True
    assert client.delete("/api/quotes/"+q["quote_id"]).json()["recoverable"]
    assert client.get("/api/cards").json()["total"] == 0
    assert client.get(f"/api/quotes/{q['quote_id']}/source").status_code == 404
    assert client.post("/api/add", files={"file": ("same.xlsx", original_bytes)}).status_code == 409
    assert client.get("/api/trash").json()["total"] == 1
    assert client.post(f"/api/quotes/{q['quote_id']}/restore").status_code == 200
    assert client.get(f"/api/quotes/{q['quote_id']}/source").content == original_bytes


@pytest.mark.parametrize("filename,data,status", [("../bad.xlsx", b"data", 400), ("bad.xls", b"data", 422), ("bad.pdf", b"not pdf", 422), ("bad.xlsx", b"PK fake", 422), ("empty.pdf", b"", 400)])
def test_bad_files(client, filename, data, status):
    assert client.post("/api/add", files={"file": (filename, data)}).status_code == status
    assert client.get("/api/cards").json()["total"] == 0


def test_size_limit(tmp_path):
    with TestClient(create_app(Settings(_env_file=None, data_dir=tmp_path, max_upload_mb=1))) as c:
        assert c.post("/api/add", files={"file": ("big.xlsx", b"0"*(1024*1024+1))}).status_code == 413


@pytest.mark.parametrize("body", [{"date_from": "2026-12-01", "date_to": "2026-01-01"}, {"amount_min": "2", "amount_max": "1"}, {"amount_min": "NaN"}, {"limit": 101}, {"offset": -1}, {"status": "fake"}])
def test_invalid_search(client, body):
    assert client.post("/api/search", json=body).status_code == 422


def test_missing_ids_and_state(client):
    for path in ("/api/quotes/missing", "/api/status/missing", "/api/jobs/missing"):
        assert client.get(path).status_code == 404
    assert client.get("/api/cards", params={"status": "unknown"}).status_code == 422
    q = upload(client)
    assert client.post("/api/status/"+q["quote_id"], json={"status": "extracting"}).status_code == 409
    result = client.post("/api/status/"+q["quote_id"], json={"status": "confirmed", "note": "원본 확인"}).json()
    assert result["status"] == "confirmed" and result["review_note"] == "원본 확인"
    assert client.patch("/api/edit/"+q["quote_id"], json={"title": "정정"}).json()["status"] == "review_required"
    assert client.post("/api/extract", json={"quote_id": q["quote_id"], "provider": "local"}).status_code == 409


def test_draft_comparison_exports(client):
    original = extract(client, upload(client)["quote_id"])
    item = {k: v for k, v in original["items"][0].items() if k not in ("item_id", "line_number")}
    item.update(quantity="12", unit_price="24000")
    draft = client.post(f"/api/quotes/{original['quote_id']}/drafts", json={"overrides": {"items": [item], "issue_date": "2026-10-17", "title": "=HYPERLINK(\"bad\")"}, "tax_rate": "10", "tax_basis": "exclusive"})
    assert draft.status_code == 201, draft.text
    draft = draft.json()
    assert draft["subtotal"] == "288000.0000" and draft["total_amount"] == "316800.0000"
    assert draft["quote_number"] is None and draft["kind"] == "draft"
    assert client.post(f"/api/quotes/{original['quote_id']}/drafts", json={"overrides": {"expected_version": 1}}).status_code == 409
    assert client.get("/api/quotes/"+original["quote_id"]).json()["items"][0]["quantity"] == "10"
    comparison = client.post("/api/compare", json={"quote_ids": [original["quote_id"], draft["quote_id"]]}).json()
    assert comparison["comparison_status"] == "conditional"
    assert comparison["pairs"][0]["price_difference"] == "-1000"
    assert client.get("/api/items/history", params={"keyword": "센서"}).json()["total"] == 2
    data = client.get(f"/api/quotes/{draft['quote_id']}/export", params={"format": "xlsx"})
    assert data.status_code == 200
    wb = load_workbook(io.BytesIO(data.content), data_only=False)
    assert not any(c.data_type == "f" for row in wb.active for c in row)
    assert any(c.value == "316800.0000" for row in wb.active for c in row)
    wb.close()
    pdf = client.get(f"/api/quotes/{draft['quote_id']}/export", params={"format": "pdf"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")
    import pdfplumber
    with pdfplumber.open(io.BytesIO(pdf.content)) as parsed:
        text = "\n".join(p.extract_text() or "" for p in parsed.pages)
    assert "온도 센서" in text and "316,800" in text


def test_local_pdf_sample(client):
    content = (Path(__file__).resolve().parents[1]/"agent_configs/samples/quote_review.pdf").read_bytes()
    q = extract(client, upload(client, content, "review.pdf")["quote_id"])
    assert len(q["items"]) == 3
    assert q["items"][2]["quantity"] is None
    assert q["items"][1]["amount"] == "190000"


def test_auth_and_cors(tmp_path):
    config = Settings(_env_file=None, data_dir=tmp_path, api_token=SecretStr("test-only-token"))
    with TestClient(create_app(config)) as c:
        assert c.get("/health").status_code == 200
        assert c.get("/api/cards").status_code == 401
        assert c.get("/api/cards", headers={"Authorization": "Bearer test-only-token"}).status_code == 200
        response = c.options("/api/cards", headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "GET"})
        assert response.headers["access-control-allow-origin"] == "http://localhost:5173"


def test_interrupted_job_recovery(tmp_path):
    config = Settings(_env_file=None, data_dir=tmp_path)
    db = Database(config)
    db.initialize()
    with db.sessions.begin() as session:
        q = Quote(id=uid(), status="extracting", extraction_status="running")
        set_fields(q, QuoteFields().model_dump(mode="json"))
        session.add(q)
        record(session, q, "test")
        j = Job(id=uid(), quote_id=q.id, kind="extract", provider="local", status="running")
        session.add(j)
        job_id = j.id
    db.engine.dispose()
    with TestClient(create_app(config)) as c:
        assert c.get("/api/jobs/"+job_id).json()["status"] == "failed"


def transport_result(result, status=200):
    calls = []
    def handle(request):
        calls.append((request.method, request.url.path))
        if request.method == "DELETE":
            return httpx.Response(200, json={"deleted": True})
        if request.url.path.endswith("/files"):
            return httpx.Response(status, json={"id": "file-test", "error": "do-not-expose-test-key"})
        if request.method == "POST":
            body = json.loads(request.content)
            assert body["config_id"] == "3" and body["include"] == ["last"]
            return httpx.Response(200, json={"id": "job-test", "status": "queued"})
        return httpx.Response(200, json={"id": "job-test", "status": "completed", "output": [{"type": "message", "content": [{"type": "output_text", "text": json.dumps(result)}]}]})
    return httpx.MockTransport(handle), calls


def test_upstage_protocol_cleanup(tmp_path):
    transport, calls = transport_result({"items": []})
    config = Settings(_env_file=None, data_dir=tmp_path, upstage_api_key="test-only-key", agent_poll_seconds=.01)
    result = UpstageClient(config, transport).run(b"%PDF-test", "agent-test", "3")
    assert result["data"] == {"items": []} and result["input_file_deleted"]
    assert calls[-1] == ("DELETE", "/v2/files/file-test")


@pytest.mark.parametrize("status", [401, 403, 429, 503])
def test_upstage_failure_sanitized(tmp_path, status):
    transport, _ = transport_result({}, status)
    config = Settings(_env_file=None, data_dir=tmp_path, upstage_api_key="test-only-key")
    with pytest.raises(AgentError) as error:
        UpstageClient(config, transport).run(b"%PDF-test", "agent-test", "3")
    assert str(status) in str(error.value)
    assert "do-not-expose-test-key" not in str(error.value)


def test_api_extraction_failure_keeps_original(client, app):
    q = upload(client)
    def fail(*args):
        raise AgentError("모의 외부 서비스 실패")
    app.state.service.agent.run = fail
    app.state.config.upstage_api_key = SecretStr("test-only-key")
    r = client.post("/api/extract", json={"quote_id": q["quote_id"], "provider": "upstage"})
    job = client.get("/api/jobs/"+r.json()["job_id"]).json()
    assert job["status"] == "failed"
    assert client.get("/api/quotes/"+q["quote_id"]).json()["status"] == "review_required"
    assert client.get(f"/api/quotes/{q['quote_id']}/source").status_code == 200


def test_tax_basis_and_discount_are_not_assumed(client):
    original = extract(client, upload(client)["quote_id"])
    url = f"/api/quotes/{original['quote_id']}/drafts"
    unspecified = client.post(url, json={"tax_rate": "10"}).json()
    assert unspecified["subtotal"] == "250000.0000"
    assert unspecified["total_amount"] is None
    inclusive = client.post(url, json={"tax_rate": "10", "tax_basis": "inclusive"}).json()
    assert inclusive["total_amount"] == "250000.0000"
    assert inclusive["subtotal"] == "227272.7273"
    assert inclusive["tax_amount"] == "22727.2727"
    discounted = client.post(url, json={"tax_rate": "10", "tax_basis": "exclusive", "overrides": {"discount_terms": "합계에서 5% 할인"}}).json()
    assert discounted["total_amount"] is None
    assert any(w["code"] == "discount_requires_review" for w in discounted["validation"]["warnings"])


def test_mixed_currency_never_summed(client):
    original = upload(client)
    rows = [{"item_name": "A", "quantity": "1", "unit_price": "10", "currency": "KRW"}, {"item_name": "B", "quantity": "1", "unit_price": "10", "currency": "USD"}]
    q = client.post(f"/api/quotes/{original['quote_id']}/drafts", json={"overrides": {"items": rows}, "tax_rate": "0"}).json()
    assert q["subtotal"] is None and q["total_amount"] is None
    assert any(w["code"] == "mixed_currencies" for w in q["validation"]["warnings"])


def test_failed_job_retry_and_busy_guard(client, app):
    q = upload(client)
    from app.schemas import ExtractionRequest
    job = app.state.service.start_extraction(ExtractionRequest(quote_id=q["quote_id"], provider="local"))
    assert client.patch("/api/edit/"+q["quote_id"], json={"title": "busy"}).status_code == 409
    assert client.delete("/api/quotes/"+q["quote_id"]).status_code == 409
    assert client.post("/api/extract", json={"quote_id": q["quote_id"], "provider": "local"}).status_code == 409
    app.state.service.fail_job(job["job_id"], "모의 중단")
    extract(client, q["quote_id"])
    assert client.post("/api/extract", json={"quote_id": q["quote_id"], "provider": "local"}).status_code == 409
    assert client.post("/api/extract", json={"quote_id": q["quote_id"], "provider": "local", "force": True}).status_code == 202


def test_compare_excludes_currency_and_units(client):
    a, b = upload(client), upload(client, xlsx(suffix="2"))
    for q, currency, unit in ((a, "KRW", "개"), (b, "USD", "세트")):
        assert client.patch("/api/edit/"+q["quote_id"], json={"currency": currency, "items": [{"item_name": "센서", "specification": "TS-100", "unit": unit, "quantity": "1", "unit_price": "10"}]}).status_code == 200
    result = client.post("/api/compare", json={"quote_ids": [a["quote_id"], b["quote_id"]]}).json()
    assert result["comparison_status"] == "insufficient_data" and len(result["excluded_rows"]) == 2
    assert client.post("/api/compare", json={"quote_ids": [a["quote_id"], a["quote_id"]]}).status_code == 422


def test_agent_workers_validate_results(client, app):
    a, b = upload(client), upload(client, xlsx(suffix="2"))
    app.state.config.upstage_api_key = SecretStr("test-only-key")
    fixtures = {
        "compare": {"comparison_status": "conditional", "comparison_rows_json": "[]", "comparable_pairs_json": "[]", "excluded_rows_json": "[]"},
        "draft": {"draft_status": "needs_input", "draft_items_json": "[]", "terms_json": "{}", "changes_json": "[]"},
    }
    def run(pdf, agent, config):
        assert pdf.startswith(b"%PDF-")
        kind = "compare" if agent == app.state.config.quote_compare_agent_id else "draft"
        return {"data": fixtures[kind], "remote_job_id": "job-mocked", "input_file_deleted": True}
    app.state.service.agent.run = run
    for kind, ids in (("compare", [a["quote_id"], b["quote_id"]]), ("draft", [a["quote_id"]])):
        response = client.post("/api/agent-runs", json={"kind": kind, "quote_ids": ids})
        assert response.status_code == 202
        result = client.get("/api/jobs/"+response.json()["job_id"]).json()
        assert result["status"] == "completed", result
    fixtures["compare"]["comparison_status"] = "bad-status"
    response = client.post("/api/agent-runs", json={"kind": "compare", "quote_ids": [a["quote_id"], b["quote_id"]]})
    assert client.get("/api/jobs/"+response.json()["job_id"]).json()["status"] == "failed"


def test_atomic_edit_validation_does_not_change_quote(client):
    q = upload(client)
    old_version = q["version"]
    response = client.patch("/api/edit/"+q["quote_id"], json={"customer_name": "실패해야 함", "items": [{"item_name": "X", "unit_price": "-1"}]})
    assert response.status_code == 422
    actual = client.get("/api/quotes/"+q["quote_id"]).json()
    assert actual["customer_name"] is None and actual["version"] == old_version


def test_upstage_timeout_and_invalid_output_cleanup(tmp_path):
    calls = []
    def handle(request):
        calls.append(request.method)
        if request.method == "DELETE":
            return httpx.Response(200, json={"deleted": True})
        if request.url.path.endswith("/files"):
            return httpx.Response(200, json={"id": "file-1"})
        return httpx.Response(200, json={"id": "job-1", "status": "in_progress"})
    config = Settings(_env_file=None, data_dir=tmp_path, upstage_api_key="test-key", agent_timeout_seconds=1, agent_poll_seconds=.01)
    with pytest.raises(AgentError, match="제한 시간"):
        UpstageClient(config, httpx.MockTransport(handle)).run(b"test", "agent", "3")
    assert calls[-1] == "DELETE"


def test_fractional_decimal_calculation(client):
    q = upload(client)
    data = client.patch("/api/edit/"+q["quote_id"], json={"items": [{"item_name": "소수 품목", "quantity": "0.1", "unit_price": "0.2", "amount": "0.02"}]}).json()
    assert data["validation"]["computed_item_amounts"] == ["0.02"]
    assert not any(w["code"] == "line_amount_mismatch" for w in data["validation"]["warnings"])
    draft = client.post(f"/api/quotes/{q['quote_id']}/drafts", json={"tax_rate": "0"}).json()
    assert draft["total_amount"] == "0.0200"
