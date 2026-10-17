import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.config import Settings
from app.documents import local_extract, make_pdf, read_document
from app.extraction import normalize_result
from app.main import create_app


def quotation():
    workbook = Workbook()
    for row in [["유효기간", "견적일로부터 7일"], ["납기", "발주 후 4주"],
                ["합 계 (VAT 별도)", 16000000],
                ["품목명", "수량", "단위", "단가", "금액"], ["장비", 1, "EA", 16000000, 16000000]]:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_relative_terms_survive_agent_date_fields_and_raw_headers():
    fields, _ = normalize_result({"valid_until": "견적일로부터 7일", "delivery_date": "발주 후 4주"})
    assert fields["validity_terms"] == "견적일로부터 7일"
    assert fields["delivery_period"] == "발주 후 4주"
    assert fields["valid_until"] is None and fields["delivery_date"] is None
    fields, _ = normalize_result({"raw_fields_json": '{"valid_until":"견적일로부터 30일","delivery_date":"일정협의"}'})
    assert fields["validity_terms"] == "견적일로부터 30일"
    assert fields["delivery_period"] == "일정협의"
    fields, _ = normalize_result({"validity_terms": "2026년10월31일까지", "delivery_date": "2026 . 11 . 01"})
    assert fields["valid_until"] == "2026-10-31" and fields["delivery_date"] == "2026-11-01"


@pytest.mark.parametrize("tax,subtotal,total", [("VAT 별도", "16000000", None), ("VAT 포함", None, "16000000")])
def test_pdf_explicit_terms_and_amount_basis(tax, subtotal, total):
    content = make_pdf("가상 견적", ["유효기간", "프로젝트명 가상 시스템 견적일로부터 7일", "(EFFECTIVE FOR)",
                                      "납기", f"견 적 금 액 일천육백만 원정 ({tax}) 발주 후 4주", "(DELIVERY)",
                                      f"합 계 ({tax}) 16,000,000"])
    fields, _ = normalize_result(local_extract(read_document(content, ".pdf")))
    assert fields["validity_terms"] == "견적일로부터 7일"
    assert fields["delivery_period"] == "발주 후 4주"
    assert fields["subtotal"] == subtotal and fields["total_amount"] == total


def test_user_direction_persists_through_extraction_edit_search_and_restart(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path, extraction_provider="local")
    content = quotation()
    app = create_app(settings)
    with TestClient(app) as client:
        uploaded = client.post("/api/add", files={"file": ("quote.xlsx", content)}, data={"direction": "sent"})
        assert uploaded.status_code == 201 and uploaded.json()["direction"] == "sent"
        quote_id = uploaded.json()["quote_id"]
        started = client.post("/api/extract", json={"quote_id": quote_id, "provider": "local"})
        assert started.status_code == 202
        assert client.get('/api/jobs/' + started.json()['job_id']).json()['status'] == 'completed'
        quote = client.get('/api/quotes/' + quote_id).json()
        assert quote['direction'] == 'sent'
        assert quote['validity_terms'] == '견적일로부터 7일'
        assert quote['delivery_period'] == '발주 후 4주'
        assert quote['quoted_amount'] == quote['subtotal'] == '16000000'
        assert quote['total_amount'] is None
        listed = client.get('/api/cards', params={'direction': 'sent'}).json()
        assert listed['total'] == 1
        assert listed['items'][0]['quoted_amount'] == '16000000'
        assert client.post('/api/search', json={'direction': 'received'}).json()['total'] == 0
        assert client.post('/api/add', files={'file': ('quote.xlsx', content)}, data={'direction': 'received'}).status_code == 409
        assert client.patch('/api/edit/' + quote_id, json={'direction': 'received', 'expected_version': quote['version']}).status_code == 200
        assert client.post('/api/extract', json={'quote_id': quote_id, 'provider': 'local', 'force': True}).status_code == 202
        assert client.get('/api/quotes/' + quote_id).json()['direction'] == 'received'
        assert client.get('/api/cards', params={'direction': 'invalid'}).status_code == 422
        assert client.patch('/api/edit/' + quote_id, json={'direction': 'invalid'}).status_code == 422
        assert client.post('/api/status/' + quote_id, json={'status': 'confirmed'}).status_code == 200
        labelled = client.patch('/api/edit/' + quote_id, json={'direction': 'received'}).json()
        assert labelled['status'] == 'confirmed'
        exported = client.get(f'/api/quotes/{quote_id}/export?format=xlsx')
        workbook = load_workbook(io.BytesIO(exported.content))
        assert any(cell.value == '발주 후 4주' for row in workbook.active for cell in row)
        workbook.close()
    with TestClient(create_app(settings)) as client:
        assert client.get('/api/quotes/' + quote_id).json()['direction'] == 'received'
        assert client.patch('/api/edit/' + quote_id, json={'direction': None}).json()['direction'] is None


def test_agent_missing_terms_recovered_and_user_label_not_inferred(tmp_path):
    app = create_app(Settings(_env_file=None, data_dir=tmp_path, upstage_api_key='test-only-key'))
    with TestClient(app) as client:
        app.state.service.agent.run = lambda *args: {'data': {'direction': 'received', 'items': []}}
        quote_id = client.post('/api/add', files={'file': ('quote.xlsx', quotation())}, data={'direction': 'sent'}).json()['quote_id']
        response = client.post('/api/extract', json={'quote_id': quote_id, 'provider': 'upstage'})
        assert client.get('/api/jobs/' + response.json()['job_id']).json()['status'] == 'completed'
        quote = client.get('/api/quotes/' + quote_id).json()
        assert quote['direction'] == 'sent'
        assert quote['delivery_period'] == '발주 후 4주'
        assert quote['validity_terms'] == '견적일로부터 7일'
        assert quote['quoted_amount'] == '16000000'
