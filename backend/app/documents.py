import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import xlrd
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .config import ROOT
from .quote_layout import pdf_line_items


class DocumentError(ValueError):
    pass


ALIASES = {
    "quote_number": ["견적번호", "quotationno", "quoteno", "견적서번호"],
    "title": ["제목", "견적명", "프로젝트명", "subject"],
    "issuer_company": ["발행회사", "공급사", "공급자", "공급업체", "issuer", "supplier"],
    "customer_name": ["고객사", "거래처", "수신회사", "수신", "受信", "customer", "client"],
    "contact_name": ["담당자", "contactname"], "contact_email": ["이메일", "email"],
    "contact_phone": ["연락처", "전화번호", "phone"],
    "issue_date": ["작성일", "발행일", "견적일", "견적일자", "發行日字", "issuedate"],
    "valid_until": ["유효일", "유효기한", "validuntil"], "delivery_date": ["납기일", "deliverydate"],
    "validity_terms": ["유효기간", "有效期間", "有効期間", "effectivefor", "validity"],
    "delivery_period": ["납기", "예정납기", "豫定納期", "delivery", "leadtime"],
    "subtotal": ["공급가액", "소계", "subtotal"], "tax_amount": ["부가세", "부가가치세", "세액", "taxamount"],
    "tax_rate": ["세율", "taxrate"], "total_amount": ["총액", "합계금액", "총금액", "grandtotal", "totalamount", "합계", "금액합계", "견적금액", "見積金額", "見積金額"],
    "tax_terms": ["세금조건"], "tax_basis": ["세금기준"],
    "currency": ["통화", "currency"], "payment_terms": ["결제조건", "지급조건", "paymentterms"],
    "delivery_terms": ["납품조건", "인도조건", "deliveryterms"],
    "discount_terms": ["할인", "할인조건", "discount"], "notes": ["비고", "특이사항", "notes"],
}
ITEM_ALIASES = {
    "item_name": ["품목명", "품명", "품목", "제품명", "品名", "itemname", "description"],
    "specification": ["규격", "사양", "모델", "specification", "model"],
    "unit": ["단위", "單位", "unit"], "quantity": ["수량", "數量", "quantity", "qty"],
    "unit_price": ["단가", "單價", "unitprice", "price"], "amount": ["금액", "품목금액", "공급가", "金額", "金額", "amount", "linetotal"],
}


def norm(value):
    return re.sub(r"[\s:：()._%\-]", "", str(value if value is not None else "")).casefold()


def validate_file(data: bytes, suffix: str, max_pages: int):
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise DocumentError("확장자와 PDF 내용이 일치하지 않습니다.")
        try:
            from pdfminer.pdfpage import PDFPage
            pages = list(PDFPage.get_pages(io.BytesIO(data), check_extractable=True))
            if not pages or len(pages) > max_pages:
                raise DocumentError(f"PDF는 1~{max_pages}페이지여야 합니다.")
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("PDF를 읽을 수 없습니다. 암호화 또는 손상 여부를 확인하세요.") from exc
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 100*1024*1024 or len(archive.infolist()) > 2000:
                    raise DocumentError("압축 해제 크기 또는 항목 수가 너무 큽니다.")
                if "xl/workbook.xml" not in archive.namelist():
                    raise DocumentError("정상적인 XLSX 파일이 아닙니다.")
                if any("vbaProject" in x.filename for x in archive.infolist()):
                    raise DocumentError("매크로 포함 문서는 지원하지 않습니다.")
        except (zipfile.BadZipFile, KeyError) as exc:
            raise DocumentError("XLSX 파일이 손상되었습니다.") from exc
    elif suffix == ".xls":
        open_xls(data).release_resources()
    else:
        raise DocumentError("현재 .xls, .xlsx와 텍스트 PDF만 지원합니다.")


def open_xls(data):
    try:
        if xlrd.inspect_format(content=data) != "xls":
            raise DocumentError("확장자와 XLS 내용이 일치하지 않습니다.")
        return xlrd.open_workbook(file_contents=data, on_demand=True, ragged_rows=True, logfile=io.StringIO())
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("XLS를 읽을 수 없습니다. 암호화 또는 손상 여부를 확인하세요.") from exc


def read_xls(data):
    workbook = open_xls(data)
    tables, raw = [], []
    warnings = ["XLS 수식・매크로・외부 연결은 실행하지 않습니다. 저장된 셀 결과만 읽으므로 계산 캐시가 없으면 원본 확인이 필요합니다."]
    try:
        if workbook.nsheets > 30:
            raise DocumentError("시트는 최대 30개까지 지원합니다.")
        for index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(index)
            if sheet.visibility != 0:
                warnings.append("숨김 시트는 견적 추출 대상에서 제외했습니다.")
                workbook.unload_sheet(index)
                continue
            if sheet.nrows * sheet.ncols > 100000:
                raise DocumentError("시트 크기가 100,000셀을 초과합니다.")
            rows, lines, character_count = [], [], 0
            for row_index in range(sheet.nrows):
                values = []
                for cell in sheet.row(row_index):
                    value = cell.value
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = ""
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode).strftime("%Y-%m-%d")
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        warnings.append("Excel 오류 셀이 있습니다. 해당 값은 추정하지 않았습니다.")
                        value = ""
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        value = format(Decimal(str(value)), "f").rstrip("0").rstrip(".") if value % 1 else str(int(value))
                    values.append(str(value).strip())
                character_count += sum(map(len, values))
                if character_count > 300000:
                    raise DocumentError("시트 텍스트가 300,000자를 초과합니다.")
                if any(values):
                    rows.append(values)
                    lines.append(f"행 {row_index+1}: " + " | ".join(f"열 {i+1}: {v}" for i, v in enumerate(values) if v))
            tables.append(rows)
            raw.append(f"[시트: {sheet.name}]\n" + "\n".join(lines))
            workbook.unload_sheet(index)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("XLS 셀을 읽지 못했습니다. 원본의 날짜・값・손상 여부를 확인하세요.") from exc
    finally:
        workbook.release_resources()
    return tables, raw, warnings


def read_document(data, suffix, max_pages=100):
    validate_file(data, suffix, max_pages)
    warnings, tables, raw, page_texts = [], [], [], []
    if suffix == ".xls":
        tables, raw, warnings = read_xls(data)
    elif suffix == ".xlsx":
        try:
            for cached in (True, False):
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=cached, keep_links=False)
                try:
                    if len(workbook.sheetnames) > 30:
                        raise DocumentError("시트는 최대 30개까지 지원합니다.")
                    for sheet in workbook:
                        if sheet.sheet_state != "visible":
                            warnings.append("숨김 시트는 견적 추출 대상에서 제외했습니다.")
                            continue
                        if sheet.max_row and sheet.max_column and sheet.max_row*sheet.max_column > 100000:
                            raise DocumentError("시트 크기가 100,000셀을 초과합니다.")
                        rows, cell_count, character_count = [], 0, 0
                        for row in sheet.iter_rows():
                            cell_count += len(row)
                            if cell_count > 100000:
                                raise DocumentError("시트 크기가 100,000셀을 초과합니다.")
                            values = []
                            for cell in row:
                                v = cell.value
                                if isinstance(v, (date, datetime)):
                                    v = v.strftime("%Y-%m-%d")
                                if not cached and cell.data_type == "f":
                                    warnings.append("Excel 수식은 실행하지 않습니다. 캐시된 결과가 없으면 값 확인이 필요합니다.")
                                values.append(str(v) if v is not None else "")
                                character_count += len(values[-1])
                                if character_count > 300000:
                                    raise DocumentError("시트 텍스트가 300,000자를 초과합니다.")
                            if any(values):
                                rows.append(values)
                        if cached:
                            tables.append(rows)
                            raw.append(f"[시트: {sheet.title}]\n" + "\n".join(" | ".join(r) for r in rows))
                        else:
                            raw.append(f"[원본 수식: {sheet.title}]\n" + "\n".join(" | ".join(r) for r in rows if any(x.startswith("=") for x in r)))
                finally:
                    workbook.close()
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("XLSX를 읽을 수 없습니다.") from exc
    else:
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for i, page in enumerate(pdf.pages):
                    page_text = page.extract_text() or ""
                    page_texts.append(page_text)
                    raw.append(f"[페이지 {i+1}]\n" + page_text)
                    tables.extend(page.extract_tables() or [])
        except ImportError:
            from pdfminer.high_level import extract_text
            raw.append(extract_text(io.BytesIO(data)))
            warnings.append("표 파서가 없어 텍스트만 읽었습니다. 품목은 직접 확인하세요.")
        except Exception as exc:
            raise DocumentError("PDF 텍스트와 표를 추출하지 못했습니다.") from exc
    text = "\n\n".join(raw)
    if len(text) > 300000:
        raise DocumentError("추출 텍스트가 300,000자를 초과합니다.")
    content_text = re.sub(r"\[(?:페이지 \d+|시트: [^\]]*|원본 수식: [^\]]*)\]", "", text)
    if not re.search(r"[가-힣A-Za-z0-9一-龥]", content_text):
        message = "텍스트가 없는 스캔 PDF입니다. OCR 지원은 별도 확장 범위입니다." if suffix == ".pdf" else "표시 시트에 읽을 수 있는 셀 내용이 없습니다."
        raise DocumentError(message)
    line_items = pdf_line_items(page_texts)
    if line_items:
        warnings.append("품목 수량·단가·금액은 명시된 PDF 텍스트 행으로 대조했습니다. 화면 표기와 텍스트 레이어가 다를 수 있으므로 원본을 확인하세요.")
    selection_required = bool(re.search(r"\d+번[^\n]*선택|선택\s*옵션|택\s*일", text))
    document = {"text": text, "tables": tables, "warnings": sorted(set(warnings)), "line_items": line_items, "selection_required": selection_required}
    if suffix == ".xls":
        candidates = local_extract(document)["items"]
        if candidates and all(all(re.fullmatch(r"\d+(?:\.\d+)?", str(row.get(k) or "").replace(",", "")) for k in ("quantity", "unit_price", "amount")) and row.get("unit") for row in candidates):
            for row in candidates:
                row["raw_fields"]["source"] = "xls_table"
            document["line_items"] = candidates
    return document


def local_extract(document):
    text, tables = document["text"], document["tables"]
    result = {"raw_fields": {}, "items": []}
    for pattern, basis in ((r"(?:VAT|부가세|부가가치세)\s*별도", "exclusive"), (r"(?:VAT|부가세|부가가치세)\s*포함", "inclusive")):
        match = re.search(pattern, text, re.I)
        if match:
            result["tax_terms"], result["tax_basis"] = match[0], basis
            break
    # Explicit label/value pairs only. Never guess missing headers from arbitrary text.
    lookup = {norm(a): k for k, aliases in ALIASES.items() for a in [k, *aliases]}
    for segment in re.split(r"\n| / |\|", text):
        match = re.match(r"\s*([^:：]+)[:：]\s*(.+)", segment)
        if match and norm(match[1]) in lookup:
            key = lookup[norm(match[1])]
            result.setdefault(key, match[2].strip())
            result["raw_fields"].setdefault(key, match[2].strip())
    for table in tables:
        mapping = None
        for row in table:
            values = [str(x if x is not None else "").strip() for x in row]
            for i, value in enumerate(values[:-1]):
                key = lookup.get(norm(re.sub(r"\([^)]*\)", "", value)))
                if not key:
                    continue
                following = next((v for v in values[i+1:] if v), "")
                if following and norm(following) not in lookup:
                    if key == "total_amount" and norm(re.sub(r"\([^)]*\)", "", value)) in {"합계", "금액합계", "견적금액", "見積金額", "見積金額"} and result.get("tax_basis") == "exclusive":
                        key = "subtotal"
                    result.setdefault(key, following)
                    result["raw_fields"].setdefault(key, following)
            header = {i: k for i, value in enumerate(values) for k, names in ITEM_ALIASES.items() if norm(value) in {norm(name) for name in names}}
            if "item_name" in header.values() and len(header) >= 2:
                if mapping is None or len(header) >= len(mapping):
                    mapping = header
                continue
            if mapping:
                item = {key: values[i] if i < len(values) else "" for i, key in mapping.items()}
                if not item.get("item_name"):
                    continue
                if norm(item["item_name"]) in lookup or norm(item["item_name"]) in {"합계", "총계", "total"}:
                    continue
                if not any(item.get(k) for k in ("quantity", "unit_price", "amount")):
                    result["raw_fields"].setdefault("unpriced_rows", []).append(item["item_name"])
                    continue
                item["original_item_name"] = item["item_name"]
                item["raw_fields"] = dict(item)
                result["items"].append(item)
    # These explicit terms can sit in a neighbouring PDF column, so the text
    # line contains the project name/amount before the term itself.
    validity = re.search(r"(?:견적(?:일|일자)?|발행일)\s*(?:로부터|부터|후)\s*\d+\s*(?:영업일|일|주|개월|달|월)", text)
    if validity:
        result.setdefault("validity_terms", validity[0])
    delivery = re.search(r"(?:발주|주문|계약|수주)\s*(?:일로부터|일?\s*후|로부터|부터)\s*\d+\s*(?:영업일|일|주|개월|달|월)(?:\s*이내)?", text)
    if delivery:
        result.setdefault("delivery_period", delivery[0])
    elif re.search(r"납\s*기\s*[:：]?\s*\n[^\n]*일정\s*협의", text):
        result.setdefault("delivery_period", "일정협의")
    number = r"(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"
    summary = rf"\s*(?:금\s*액\s*합\s*계|합\s*계|총\s*액|총\s*계|견\s*적\s*금\s*액)\s*(?:\([^)]*\))?\s*[:：]?\s*(?:KRW|₩)?\s*(?P<amount>{number})\s*원?\s*(?:\([^)]*\))?\s*"
    for line in text.splitlines():
        match = re.fullmatch(summary, line, re.I)
        if match:
            key = "subtotal" if result.get("tax_basis") == "exclusive" else "total_amount"
            result.setdefault(key, match["amount"])
            result["raw_fields"].setdefault(key, match["amount"])
    if document.get("line_items"):
        result["items"] = document["line_items"]
    if document.get("selection_required"):
        result["raw_fields"]["selection_required"] = True
    if result.get("tax_amount"):
        tax = re.fullmatch(r"([\d,]+(?:\.\d+)?)\s*원?\s*\(([\d.]+)\s*%\)", str(result["tax_amount"]))
        if tax:
            result["tax_amount"] = tax[1]
            result.setdefault("tax_rate", tax[2])
    return result


def font_name(custom=None):
    name = "QuoteFlowKorean"
    if name not in pdfmetrics.getRegisteredFontNames():
        path = custom or ROOT / "app" / "fonts" / "NanumGothic-Regular.ttf"
        pdfmetrics.registerFont(TTFont(name, str(path)))
    return name


def make_pdf(title, lines, custom_font=None):
    output = io.BytesIO()
    font = font_name(custom_font)
    style = ParagraphStyle("body", fontName=font, fontSize=9, leading=14, wordWrap="CJK")
    parts = [Paragraph(escape(title), ParagraphStyle("title", parent=style, fontSize=16, leading=23)), Spacer(1, 14)]
    for line in lines:
        # Keep long JSON and user strings breakable and escape ReportLab markup.
        for start in range(0, len(str(line)), 1500):
            parts.append(Paragraph(escape(str(line)[start:start+1500]) or " ", style))
            parts.append(Spacer(1, 5))
    SimpleDocTemplate(output, title=title, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36).build(parts, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()


def quote_pdf(quote, custom_font=None):
    output = io.BytesIO()
    font = font_name(custom_font)
    style = ParagraphStyle("body", fontName=font, fontSize=9, leading=13, wordWrap="CJK")
    p = lambda value: Paragraph(escape(str(value if value is not None else "미기재")), style)
    title_style = ParagraphStyle("title", parent=style, fontSize=18, leading=25)
    parts = [Paragraph("견적 아카이브 사본", title_style), Spacer(1, 14)]
    parts.append(p(f"견적서 구분: {DIRECTION_LABELS.get(quote.get('direction'), '미분류')}"))
    for key, label in [("quote_number", "견적번호"), ("title", "제목"), ("issuer_company", "발행회사"), ("customer_name", "고객사"), ("issue_date", "작성일"), ("valid_until", "유효기간"), ("delivery_date", "납기일")]:
        parts.append(p(f"{label}: {quote.get(key) or '미기재'}"))
    parts.append(Spacer(1, 14))
    keys = ["item_name", "specification", "unit", "quantity", "unit_price", "amount"]
    rows = [[p(x) for x in ["품목명", "규격", "단위", "수량", "단가", "금액"]]]
    for item in quote.get("items", []):
        rows.append([p(display_number(item.get(k)) if k in ("quantity", "unit_price", "amount") else item.get(k)) for k in keys])
    table = Table(rows, colWidths=[100, 145, 35, 55, 85, 100], repeatRows=1, splitInRow=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef8")), ("GRID", (0, 0), (-1, -1), .3, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    parts += [table, Spacer(1, 14)]
    for key, label in EXPORT_LABELS.items():
        value = quote.get(key)
        if key in ("subtotal", "tax_rate", "tax_amount", "total_amount"):
            value = display_number(value)
        parts.append(p(f"{label}: {value if value is not None else '미기재'}"))
    parts += [Spacer(1, 10), p(f"상태: {STATUS_LABELS[quote['status']]} / 버전: {quote['version']} / 담당자 원본 대조 필요")]
    SimpleDocTemplate(output, leftMargin=22, rightMargin=22, topMargin=30, bottomMargin=30).build(parts, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()


def quote_xlsx(quote):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "견적서"
    rows = [["견적 아카이브 사본"], ["상태", STATUS_LABELS[quote["status"]]], ["버전", quote["version"]]]
    rows.append(["견적서 구분", DIRECTION_LABELS.get(quote.get("direction"), "미분류")])
    for key in ("quote_number", "title", "issuer_company", "customer_name", "issue_date", "valid_until", "delivery_date", "currency"):
        label = {"quote_number": "견적번호", "title": "제목", "issuer_company": "발행 회사", "customer_name": "고객사", "issue_date": "작성일", "valid_until": "유효기간", "delivery_date": "납기일", "currency": "통화"}[key]
        rows.append([label, quote.get(key)])
    rows += [[], ["품목명", "규격", "단위", "수량", "단가", "금액"]]
    rows += [[item.get(k) for k in ("item_name", "specification", "unit", "quantity", "unit_price", "amount")] for item in quote.get("items", [])]
    rows.append([])
    for key, label in EXPORT_LABELS.items():
        if key != "currency":
            rows.append([label, quote.get(key)])
    rows.append(["안내", "금액은 정밀도 보존용 문자열입니다. 담당자 검토 후 사용하세요."])
    for row in rows:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Never execute user-supplied spreadsheet formulas.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in "ABCDEF":
        sheet.column_dimensions[col].width = 24 if col in "AB" else 18
    sheet["A1"].font = Font(size=16, bold=True)
    sheet["A1"].fill = PatternFill("solid", fgColor="E8EEF8")
    sheet.freeze_panes = "A14"
    sheet.print_options.horizontalCentered = True
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def agent_input(title, data, custom_font=None):
    return make_pdf(title, json.dumps(data, ensure_ascii=False, indent=2, default=str).splitlines(), custom_font)


EXPORT_LABELS = {"currency": "통화", "subtotal": "공급가액", "tax_rate": "세율 (%)", "tax_terms": "세금 조건", "tax_amount": "부가세", "total_amount": "총액", "validity_terms": "유효기간 조건", "delivery_period": "납기 조건", "payment_terms": "결제 조건", "delivery_terms": "납품 조건", "discount_terms": "할인 조건", "notes": "비고"}
STATUS_LABELS = {"review_required": "확인 필요", "extracting": "추출 중", "confirmed": "검토 완료", "rejected": "반려", "archived": "보관"}
DIRECTION_LABELS = {"sent": "보낸 견적서", "received": "받은 견적서"}


def display_number(value):
    if value is None:
        return "미기재"
    text = format(Decimal(str(value)), ",f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def page_number(canvas, document):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(document.pagesize[0]-document.rightMargin, 14, str(document.page))
    canvas.restoreState()
