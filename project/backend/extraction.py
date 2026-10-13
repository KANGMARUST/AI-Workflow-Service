from __future__ import annotations
import re, os, json
from pathlib import Path
from datetime import date
import domain as d

ROW=re.compile(r'^(?:(\d+)\s+)?(.+?)\s+([\d,.]+)\s+([A-Za-z가-힣]+)\s+([\d,.]+)\s+([\d,.]+)(?:\s.*)?$')
OCR_ENGINE=None

def ocr(image):
    global OCR_ENGINE
    if OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR
        model=Path(os.getenv('QUOTE_OCR_MODEL') or d.BASE/'models'/'korean_rec.onnx')
        OCR_ENGINE=RapidOCR(rec_model_path=str(model)) if model.is_file() else RapidOCR()
    result,_=OCR_ENGINE(image)
    grouped=[]
    for box,text,_confidence in sorted(result or [],key=lambda row:sum(point[1] for point in row[0])/4):
        cy=sum(point[1] for point in box)/4;cx=min(point[0] for point in box);height=max(point[1] for point in box)-min(point[1] for point in box)
        found=next((line for line in grouped if abs(line[0]-cy)<max(height*.55,10)),None)
        if found is None:grouped.append([cy,[(cx,text)]])
        else:found[1].append((cx,text))
    return '\n'.join(' '.join(text for _,text in sorted(words)) for _,words in grouped)

def sanitize_pdf(text):
    # Some customer files cover contact text visually but retain hidden PDF text.
    # Never publish invisible contacts reconstructed from the underlying text layer.
    text=re.sub(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', '[이메일 원본 확인]',text)
    text=re.sub(r'(?:\+82[- ]?)?0\d{1,2}[- )]*\d{3,4}[- ]*\d{4}', '[연락처 원본 확인]',text)
    text=re.sub(r'^.*(?:\[연락처 원본 확인\]|\[이메일 원본 확인\]).*$', '[담당자·연락처는 원본 화면에서 확인]',text,flags=re.M)
    text=re.sub(r'^.*(?:대표\s*이사|代\s*表\s*理\s*事|代\s*表\s*理\s*事).*$','[대표자 정보는 원본 화면에서 확인]',text,flags=re.M)
    return text

def item(name,qty,unit,price,amount,line,component=False,parent=None):
    return dict(name=name.strip(),specification='',quantity=d.n(d.decimal(qty)),unit=unit,unit_price=d.n(d.decimal(price)),amount=d.n(d.decimal(amount)),original_text=line,is_component=component,parent_index=parent,option_group='')

def header(text):
    q={};lines=text.splitlines()
    customer=re.search(r'(?:수\s*신|受\s*信|거래처|고객사|Customer)\s*[:：]?\s*(.+)',text,re.I)
    if customer:q['customer']=re.split(r'견적번호|貴下|견적일',customer.group(1))[0].strip()
    reference=re.search(r'(?:참\s*조|參\s*照)\s*[:：]?\s*(.+)',text)
    if reference:q['customer_contact']=re.split(r'귀하|貴下',reference.group(1))[0].strip()
    for key,label in [('contact','담당자'),('phone','연락처|전화'),('email','이메일|E-?MAIL')]:
        value=re.search(r'^(?:'+label+r')\s*[:：]\s*(.+)$',text,re.M|re.I)
        if value:q[key]=value.group(1).strip()
    dt=re.search(r'(20\d{2})\s*[년./-]\s*(\d{1,2})\s*[월./-]\s*(\d{1,2})',text)
    if dt:
        try:q['issue_date']=date(*map(int,dt.groups())).isoformat()
        except ValueError:pass
    num=re.search(r'견적번호\s*[:：]?\s*([\w-]+)',text)
    if num:q['number']=num.group(1)
    project=re.search(r'(?:프로젝트명|프로젝트|題\s*目)\s*[:：]?\s*(.+)',text)
    if project:q['project_name']=re.split(r'견적일로부터',project.group(1))[0].strip()
    valid=re.search(r'견적일로부터\s*\d+\s*일',text)
    if valid:q['valid_until']=valid.group()
    for key,label in [('delivery','납기'),('payment_terms','결제조건'),('ship_to','납품장소')]:
        for idx,line in enumerate(lines):
            if line.strip()==label and idx+1<len(lines):q[key]=lines[idx+1].strip();break
    if '엔시스템' in text:q['issuer']='엔시스템'
    if q.get('delivery') and 'VAT별도)' in q['delivery']:q['delivery']=q['delivery'].split('VAT별도)',1)[-1].strip()
    if q.get('ship_to'):
        tail=re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\s+(.+)',q['ship_to'])
        if tail:q['ship_to']=tail.group(1).strip()
        elif '[담당자' in q['ship_to']:q['ship_to']=''
    q['tax_rate']=None;q['tax']=None;q['subtotal']=None;q['total']=None
    q['tax_note']='VAT 별도' if re.search(r'VAT\s*별도|부가세\s*별도',text,re.I) else 'VAT 포함' if re.search(r'VAT\s*포함|부가세\s*포함',text,re.I) else '원본 미기재'
    tax=re.search(r'(?:VAT|부가세|세율)\s*[:：]?\s*(\d+(?:\.\d+)?)\s*%',text,re.I)
    if tax:q['tax_rate']=float(tax.group(1))
    total=re.findall(r'(?:합\s*계|총\s*액|금\s*액\s*합\s*계).*?([\d,]{3,}(?:\.\d+)?)',text)
    if total:
        q['source_total']=d.n(d.decimal(total[-1]));q['source_total_kind']='subtotal' if q['tax_note']=='VAT 별도' else 'total'
        q['subtotal' if q['source_total_kind']=='subtotal' else 'total']=q['source_total']
    for key,label in [('subtotal','공급가액'),('tax','부가세'),('total','총액')]:
        value=re.search(label+r'\s*[:：]?\s*([\d,]+(?:\.\d+)?)',text)
        if value:q[key]=d.n(d.decimal(value.group(1)))
    return q

def parse_lines(text):
    items=[];indexes={};parent=None
    for line in text.splitlines():
        if '이하여백' in re.sub(r'\s+','',line):continue
        match=ROW.match(line.strip())
        if match:
            index,name,qty,unit,price,amount=match.groups()
            if name.strip() in ('합계','총액') or '합 계' in name:continue
            if index:
                signature=(index,name.strip(),qty,price,amount)
                if signature in indexes:parent=indexes[signature];continue
                parent=len(items);indexes[signature]=parent
                items.append(item(name,qty,unit,price,amount,line))
            else:
                # Unnumbered detail rows underneath packages are components.
                component=parent is not None and bool(indexes)
                items.append(item(name,qty,unit,price,amount,line,component,parent if component else None))
        elif items and re.match(r'^\s*[-•](?:\.|\s)',line):
            items[-1]['specification']+=(('\n' if items[-1]['specification'] else '')+line.strip())
    if not items:
        alt=re.compile(r'^(.+?)\s+([A-Za-z가-힣]{1,8})\s*([\d,.]+)\s+([\d,.]+)\s+([\d,.]+)$')
        for line in text.splitlines():
            m=alt.match(line.strip())
            if m:
                name,unit,qty,price,amount=m.groups();entry=item(name,qty,unit,price,amount,line)
                if '규격' in text and len(name.split())>1:
                    parts=name.rsplit(' ',1);entry['name']=parts[0];entry['specification']=parts[1]
                items.append(entry)
    return items

def read_spreadsheet(path=None,sheets=None):
    if sheets is not None:return parse_sheets(sheets)
    sheets=[]
    if path.suffix.lower()=='.xls':
        import xlrd
        book=xlrd.open_workbook(path)
        for sheet in book.sheets():sheets.append((sheet.name,[sheet.row_values(i) for i in range(sheet.nrows)]))
    else:
        from openpyxl import load_workbook
        book=load_workbook(path,data_only=True)
        for sheet in book:sheets.append((sheet.title,[[c.value for c in row] for row in sheet]))
    return parse_sheets(sheets)

def parse_sheets(sheets):
    texts=[];items=[];q={};extra=[];option_indexes=[]
    labels={'customer_contact':['참조','參照'],'contact':['견적담당자','담당자'],'phone':['연락처','전화'],'email':['이메일','email'],'customer':['수신','受信','거래처','고객사','customer'],'number':['견적번호','number'],'issue_date':['發行日字','발행일자','견적일자','작성일','issuedate'],'project_name':['題目','프로젝트명','projectname'],'delivery':['豫定納期','납기','delivery'],'ship_to':['納品場所','납품장소','shipto'],'payment_terms':['支拂條件','결제조건','paymentterms'],'valid_until':['有效期間','유효기간','validuntil'],'currency':['통화','currency'],'tax_rate':['세율','부가세율','taxrate'],'discount':['할인','discount']}
    for sheet_name,rows in sheets:
        if sheet_name.lower()=='wdata':continue
        texts.append('시트: '+sheet_name)
        header_row=None;columns={}
        for row_index,row in enumerate(rows):
            vals=[str(v).strip() if v is not None else '' for v in row]
            line=' '.join(v for v in vals if v);texts.append(line)
            normalized=[re.sub(r'\s+','',v).lower() for v in vals]
            for key,alias in labels.items():
                for j,val in enumerate(normalized):
                    if val in [x.lower() for x in alias]:
                        scan=vals[j+1:j+7] if key=='customer_contact' else vals[j+1:]
                        nextval=next((v for v in scan if v and v not in ('貴下',)),None)
                        if key=='customer_contact' and nextval is None and row_index>0:
                            prior=rows[row_index-1]
                            candidate=prior[j+2] if len(prior)>j+2 else None
                            if candidate and isinstance(candidate,str):nextval=candidate
                        if nextval and key not in q:q[key]=nextval
            if '선택 옵션' in line or '중 선택' in line or '선택옵션' in line:
                option_indexes += re.findall(r'(\d+)번',line)
                extra.append('대체 옵션이 포함된 원본입니다. 모든 옵션과 원본의 선택 조건을 보존합니다.')
            # Column meanings differ per company; use labels, not fixed coordinates.
            map_alias={'name':['품명','品名','품목명','description','name'],'quantity':['수량','數量','qty','quantity'],'unit':['단위','單位','unit'],'unit_price':['단가','單價','unitprice'],'amount':['금액','金額','공급가','amount'],'specification':['규격','사양','specification']}
            found={}
            for key,alias in map_alias.items():
                for j,val in enumerate(normalized):
                    if val in [x.lower() for x in alias]:found[key]=j;break
            if 'name' in found and ('quantity' in found or 'amount' in found):columns=found;header_row=True;continue
            if header_row and columns:
                name=vals[columns['name']] if columns['name']<len(vals) else ''
                def value(k):
                    j=columns.get(k)
                    return row[j] if j is not None and j<len(row) else None
                try:
                    qty=d.decimal(value('quantity'));price=d.decimal(value('unit_price'));amt=d.decimal(value('amount'))
                except ValueError:continue
                if name and name.lower() not in ('(description)','description','합계','합 계') and any(v is not None for v in (qty,price,amt)):
                    entry=item(name,qty, str(value('unit') or ''),price,amt,line);entry['specification']=str(value('specification') or '')
                    if vals[0].removesuffix('.0') in option_indexes:entry.update(option_group='선택1')
                    items.append(entry)
    text='\n'.join(texts);generic=header(text);generic.update(q)
    if isinstance(generic.get('issue_date'),str):
        dt=re.search(r'(20\d{2})\s*[년./-]\s*(\d{1,2})\s*[월./-]\s*(\d{1,2})',generic['issue_date'])
        if dt:
            try:generic['issue_date']=date(*map(int,dt.groups())).isoformat()
            except ValueError:generic['issue_date']=''
    generic['items']=items or parse_lines(text);generic['extraction_warnings']=extra
    return generic,text

def external_structure(text):
    if os.getenv('EXTERNAL_AI_ENABLED','false').lower()!='true' or not os.getenv('OPENAI_API_KEY'):return None
    import httpx
    prompt='이 견적 텍스트를 JSON object로 추출하세요. 필드 customer,issue_date,number,items(name,specification,unit,quantity,unit_price,amount,is_component,parent_index),source_total. 없는 값은 null. 반복 패키지는 한 번만 넣고 부품은 상위행 0 기반 인덱스 지정. 추측 금지.'
    result=httpx.post(os.getenv('OPENAI_BASE_URL','https://api.openai.com/v1').rstrip('/')+'/chat/completions',headers={'Authorization':'Bearer '+os.environ['OPENAI_API_KEY']},json={'model':os.getenv('OPENAI_MODEL','gpt-4.1-mini'),'messages':[{'role':'system','content':prompt},{'role':'user','content':text}],'response_format':{'type':'json_object'}},timeout=60)
    result.raise_for_status();return json.loads(result.json()['choices'][0]['message']['content'])

def extract(path,allow_external=True):
    suffix=path.suffix.lower();method='';extra=[]
    if suffix in ('.xls','.xlsx'):
        q,text=read_spreadsheet(path);method='local-spreadsheet'
    elif suffix=='.pdf':
        import pdfplumber
        parts=[];ocr_used=False
        with pdfplumber.open(path) as book:
            for page in book.pages:
                text=page.extract_text() or ''
                if len(text.strip())<30:text=ocr(page.to_image(resolution=180).original);ocr_used=True
                parts.append(text)
        text=sanitize_pdf('\n'.join(parts));q=header('\n'.join(parts));q['items']=parse_lines(text)
        q.update(contact='',phone='',email='')
        method='local-pdf+ocr' if ocr_used else 'local-pdf'
        extra.append('담당자·연락처는 원본 화면에서 확인하고 입력해주세요. PDF 숨김 연락처 텍스트는 표시하지 않습니다.')
    elif suffix=='.docx':
        from docx import Document
        doc=Document(path);lines=[p.text for p in doc.paragraphs]
        for table in doc.tables:
            lines.extend(' '.join(c.text for c in row.cells) for row in table.rows)
        text='\n'.join(lines);q=header(text);q['items']=parse_lines(text)
        tableq,_=parse_sheets([(str(i),[[c.text for c in row.cells] for row in table.rows]) for i,table in enumerate(doc.tables)])
        if tableq['items']:q['items']=tableq['items']
        method='local-docx'
    elif suffix in ('.png','.jpg','.jpeg'):
        text=ocr(str(path));q=header(text);q['items']=parse_lines(text);method='local-ocr';extra.append('OCR 추출 결과를 원본 이미지와 비교해주세요.')
    else:raise ValueError('지원 파일: PDF, XLS, XLSX, DOCX, PNG, JPG')
    remote=external_structure(text) if allow_external else None
    if remote:q.update(remote);method+=' + external-ai'
    q.update(source_text=sanitize_pdf(text),extraction_method=method)
    q['source_snapshot_version']=2
    q['original_fields']=json.loads(json.dumps({key:q.get(key) for key in ('number','customer','contact','customer_contact','phone','email','project_name','issue_date','subtotal','tax','tax_rate','total','source_total','source_total_kind','tax_note','items')}))
    q['extraction_warnings']=(q.get('extraction_warnings') or [])+extra
    if not q.get('items'):q['extraction_warnings'].append('인식 가능한 품목 행을 찾지 못했습니다. 품목을 직접 추가해주세요.')
    return d.empty_quote(q)
