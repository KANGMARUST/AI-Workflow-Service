"""Existing quote archive tests; no quotation creation, repricing or export workflow."""
from __future__ import annotations

import copy
from contextlib import closing
import hashlib
import io
import json
import os
from pathlib import Path
import re
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
import uuid

ROOT = Path(__file__).resolve().parents[1]
PACKAGES = ROOT / 'backend' / '.packages'
if PACKAGES.exists():
    sys.path.insert(0, str(PACKAGES))


class QuoteIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.run_root = (ROOT / 'tests' / '.runs').resolve()
        cls.run_root.mkdir(exist_ok=True)
        cls.temporary = tempfile.TemporaryDirectory(prefix='archive-integration-', dir=cls.run_root)
        cls.work = Path(cls.temporary.name)
        cls.samples=cls.work/'samples'
        cls.samples.mkdir()
        sources=list((ROOT/'data').glob('*.pdf'))+list((ROOT/'data').glob('*.xls'))
        if not sources:
            # The user may replace data/Input_data. Recover regression sources read-only
            # from the already archived files, without restoring their working folders.
            live_database=ROOT/'backend'/'storage'/'quotes.db'
            with closing(sqlite3.connect(live_database.as_uri()+'?mode=ro',uri=True)) as connection:
                for (payload,) in connection.execute('SELECT payload FROM quotes'):
                    q=json.loads(payload);name=q.get('source_name','')
                    if name.startswith(('18KJ_ACT','26KJ_POSCODX','26KJ_포스코디엑스')) and not q.get('parent_id'):
                        source=live_database.parent/'files'/Path(q['source_url']).name
                        if source.is_file():(cls.samples/name).write_bytes(source.read_bytes())
        else:
            for source in sources:(cls.samples/source.name).write_bytes(source.read_bytes())
        cls.token = uuid.uuid4().hex[:8]
        cls.counter = 0
        with socket.socket() as listener:
            listener.bind(('127.0.0.1', 0))
            cls.port = listener.getsockname()[1]
        cls.base = f'http://127.0.0.1:{cls.port}'
        cls.database = cls.work / 'storage' / 'quotes.db'
        env = os.environ.copy()
        env['QUOTE_DATA_DIR'] = str(cls.work / 'storage')
        env['QUOTE_SAMPLE_DIR'] = str(cls.samples)
        env['DATABASE_URL'] = 'sqlite:///' + cls.database.as_posix()
        env['EXTERNAL_AI_ENABLED'] = 'false'
        env.pop('OPENAI_API_KEY', None)
        env['PYTHONPATH'] = os.pathsep.join([str(PACKAGES), str(ROOT / 'backend'), env.get('PYTHONPATH', '')])
        cls.environment = env
        cls.log = open(cls.work / 'server.log', 'w', encoding='utf-8')
        cls.server = cls.start_server()
        cls.addClassCleanup(cls.cleanup)
        cls.wait_ready()
        cls.make_fixtures()

    @classmethod
    def start_server(cls):
        return subprocess.Popen([sys.executable, '-m', 'uvicorn', 'app:app', '--host', '127.0.0.1', '--port', str(cls.port)], cwd=ROOT / 'backend', env=cls.environment, stdout=cls.log, stderr=subprocess.STDOUT, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))

    @classmethod
    def wait_ready(cls):
        deadline = time.monotonic() + 45
        while time.monotonic() < deadline:
            try:
                status, health, _ = cls.request('GET', '/api/health')
                if status == 200 and health.get('ok'):
                    return
            except (OSError, ValueError):
                pass
            if cls.server.poll() is not None:
                break
            time.sleep(0.2)
        cls.log.flush()
        raise RuntimeError('Isolated API did not become ready: ' + (cls.work / 'server.log').read_text(encoding='utf-8')[-6000:])

    @classmethod
    def cleanup(cls):
        if getattr(cls, 'server', None) and cls.server.poll() is None:
            cls.server.terminate()
            try:
                cls.server.wait(timeout=10)
            except subprocess.TimeoutExpired:
                cls.server.kill()
                cls.server.wait(timeout=5)
        if getattr(cls, 'log', None):
            cls.log.close()
        if getattr(cls, 'temporary', None):
            if not cls.work.resolve().is_relative_to(cls.run_root):
                raise RuntimeError('Refusing cleanup outside tests/.runs')
            cls.temporary.cleanup()

    @classmethod
    def request(cls, method, route, data=None, body=None, content_type=None):
        headers = {}
        if data is not None:
            body = json.dumps(data, ensure_ascii=False).encode('utf-8')
            headers['Content-Type'] = 'application/json'
        elif content_type:
            headers['Content-Type'] = content_type
        req = Request(cls.base + route, data=body, method=method, headers=headers)
        try:
            response = urlopen(req, timeout=60)
        except HTTPError as exc:
            response = exc
        with response:
            raw = response.read()
            content = response.headers.get('Content-Type', '')
            parsed = json.loads(raw) if 'json' in content and raw else raw
            return response.status, parsed, {k.lower(): v for k, v in response.headers.items()}

    def api(self, method, route, data=None, expected=200):
        status, result, _ = self.request(method, route, data=data)
        self.assertEqual(status, expected, f'{method} {route}: {result}')
        return result

    @classmethod
    def multipart(cls, files):
        boundary = 'ArchiveTest' + uuid.uuid4().hex
        parts = []
        for filename, data in files:
            parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="{filename}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode('utf-8') + data + b'\r\n')
        parts.append(f'--{boundary}--\r\n'.encode('ascii'))
        return b''.join(parts), f'multipart/form-data; boundary={boundary}'

    def upload(self, paths):
        body, content = self.multipart([(p.name, p.read_bytes()) for p in paths])
        status, result, _ = self.request('POST', '/api/upload', body=body, content_type=content)
        self.assertEqual(status, 200, str(result))
        return result

    @classmethod
    def make_fixtures(cls):
        from openpyxl import Workbook
        from docx import Document
        from PIL import Image, ImageDraw, ImageFont
        rows = [
            ['견적서'], ['견적번호', 'FIXTURE-' + cls.token], ['거래처', '테스트 고객'],
            ['작성일', '2026-10-13'], ['통화', 'KRW'], ['품목명', '규격', '단위', '수량', '단가', '금액'],
            ['브라켓', '100mm', 'EA', 3, 12000, 36000], ['볼트', 'M8', 'EA', 5, 1000, 5000],
            ['공급가액', 41000], ['부가세', 4100], ['총액', 45100],
        ]
        book = Workbook()
        for row in rows:
            book.active.append(row)
        cls.xlsx = cls.work / 'fixture.xlsx'
        book.save(cls.xlsx)
        document = Document()
        for row in rows[:5]:
            document.add_paragraph(' : '.join(map(str, row)))
        table = document.add_table(rows=1, cols=6)
        for cell, label in zip(table.rows[0].cells, rows[5]):
            cell.text = label
        for row in rows[6:8]:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = str(value)
        for row in rows[8:]:
            document.add_paragraph(' : '.join(map(str, row)))
        cls.docx = cls.work / 'fixture.docx'
        document.save(cls.docx)
        picture = Image.new('RGB', (1800, 1000), 'white')
        draw = ImageDraw.Draw(picture)
        font = ImageFont.truetype('C:/Windows/Fonts/malgun.ttf', 38)
        lines = ['견적서', '견적번호 : OCR-' + cls.token, '거래처 : 테스트 고객', '작성일 : 2026-10-13',
                 '품목명    규격    단위    수량    단가    금액', '브라켓    100mm    EA    3    12000    36000',
                 '공급가액 : 36000', '부가세 : 3600', '총액 : 39600']
        for i, line in enumerate(lines):
            draw.text((60, 35 + i * 90), line, fill='black', font=font)
        cls.png = cls.work / 'fixture.png'
        picture.save(cls.png)
        cls.invalid = cls.work / 'unsupported.txt'
        cls.invalid.write_text('Unsupported file.', encoding='utf-8')
        cls.empty = cls.work / 'empty.pdf'
        cls.empty.write_bytes(b'')

    def new_archive(self):
        from openpyxl import load_workbook
        type(self).counter += 1
        book = load_workbook(self.xlsx)
        book.active['B2'] = f'ARCHIVE-{self.token}-{self.counter}'
        path = self.work / f'archive-{self.counter}.xlsx'
        book.save(path)
        result = self.upload([path])
        self.assertEqual(result['errors'], [])
        return result['quotes'][0]

    def confirm(self, quote):
        return self.api('POST', f'/api/quotes/{quote["id"]}/confirm', {'acknowledged_warnings': True})

    def test_01_real_documents_preserve_original_amounts_all_items_and_options(self):
        imported = self.api('POST', '/api/import-samples')
        self.assertEqual(imported['errors'], [])
        self.assertEqual(len(imported['quotes']), 3)
        pdfs = [q for q in imported['quotes'] if q['source_name'].lower().endswith('.pdf')]
        amounts = {float(q['subtotal']): q for q in pdfs}
        self.assertEqual(set(amounts), {16000000, 266740000})
        self.assertEqual(len(amounts[16000000]['items']), 3)
        complex_quote = amounts[266740000]
        self.assertEqual(sum(not i['is_component'] for i in complex_quote['items']), 5)
        components = [i for i in complex_quote['items'] if i['is_component']]
        self.assertEqual(len(components), 15)
        for item in components:
            self.assertIsInstance(item['parent_index'], int)
            self.assertFalse(complex_quote['items'][item['parent_index']]['is_component'])
        for quote in pdfs:
            self.assertIsNone(quote['tax_rate'])
            self.assertIsNone(quote['tax'])
            self.assertIsNone(quote['total'])
            self.assertEqual(float(quote['source_total']), float(quote['subtotal']))
            self.assertEqual(quote['source_total_kind'], 'subtotal')
            self.assertIsNone(re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', quote['source_text']))
            self.assertIsNone(re.search(r'01[016789][- ]?\d{3,4}[- ]?\d{4}', quote['source_text']))
        xls = next(q for q in imported['quotes'] if q['source_name'].lower().endswith('.xls'))
        self.assertEqual(len(xls['items']), 3)
        self.assertEqual({float(i['amount']) for i in xls['items']}, {6750000, 2910000, 2320000})
        self.assertIsNone(xls['subtotal'])
        self.assertIsNone(xls['total'])
        self.assertIsNone(xls['source_total'])
        self.assertIsNone(xls['tax'])
        options = [i for i in xls['items'] if i.get('option_group')]
        self.assertEqual(len(options), 2)
        self.assertEqual(len({i['option_group'] for i in options}), 1)
        # Options are original document information; choosing one is not required.
        reviewed = self.confirm(xls)
        self.assertEqual(reviewed['review_status'], 'confirmed')
        self.assertEqual(len(reviewed['items']), 3)
        self.assertIsNone(reviewed['subtotal'])
        duplicate = self.api('POST', '/api/import-samples')
        self.assertEqual({q['id'] for q in imported['quotes']}, {q['id'] for q in duplicate['quotes']})

    def test_02_source_files_are_linked_and_unmodified(self):
        imported = self.api('POST', '/api/import-samples')
        for quote in imported['quotes']:
            raw = (self.samples / quote['source_name']).read_bytes()
            status, stored, _ = self.request('GET', quote['source_url'])
            self.assertEqual(status, 200)
            self.assertEqual(hashlib.sha256(raw).digest(), hashlib.sha256(stored).digest())
            detail = self.api('GET', f'/api/quotes/{quote["id"]}')
            self.assertEqual(detail['source_url'], quote['source_url'])
            self.assertTrue(detail['source_text'])

    def test_03_xlsx_docx_import_and_partial_failures(self):
        result = self.upload([self.xlsx, self.docx, self.invalid, self.empty])
        self.assertEqual(len(result['quotes']), 2)
        self.assertEqual(len(result['errors']), 2)
        for quote in result['quotes']:
            self.assertEqual(len(quote['items']), 2)
            self.assertEqual(float(quote['subtotal']), 41000)
            self.assertEqual(float(quote['tax']), 4100)
            self.assertEqual(float(quote['total']), 45100)
            self.assertIsNone(quote['tax_rate'], 'A rate is absent from these source documents')
            self.assertEqual(quote['customer'], '테스트 고객')
        duplicate = self.upload([self.xlsx])
        self.assertEqual(duplicate['quotes'][0]['id'], result['quotes'][0]['id'])

    def test_04_local_korean_ocr_preserves_recognized_item_fields(self):
        self.assertTrue(self.api('GET', '/api/health')['ocr_available'])
        result = self.upload([self.png])
        self.assertEqual(result['errors'], [])
        quote = result['quotes'][0]
        self.assertIn('ocr', quote['extraction_method'].lower())
        bracket = next((i for i in quote['items'] if '브라켓' in i['name']), None)
        self.assertIsNotNone(bracket, quote['source_text'])
        self.assertIn('100mm', bracket['specification'])
        self.assertEqual(float(bracket['quantity']), 3)
        self.assertEqual(float(bracket['unit_price']), 12000)
        self.assertEqual(float(bracket['amount']), 36000)

    def test_05_correction_does_not_reprice_or_overwrite_original_amounts(self):
        quote = self.new_archive()
        original = copy.deepcopy(quote)
        quote['items'][0]['quantity'] = 8
        quote['source_url'] = '/api/files/forged.pdf'
        quote['review_status'] = 'confirmed'
        corrected = self.api('PUT', f'/api/quotes/{quote["id"]}', quote)
        self.assertEqual(float(corrected['items'][0]['amount']), 36000)
        self.assertEqual(float(corrected['subtotal']), 41000)
        self.assertEqual(float(corrected['tax']), 4100)
        self.assertEqual(float(corrected['total']), 45100)
        self.assertEqual(float(corrected['items'][0]['computed_amount']), 96000)
        self.assertTrue(corrected['warnings'])
        self.assertEqual(corrected['source_url'], original['source_url'])
        self.assertEqual(corrected['review_status'], 'pending')
        history = self.api('GET', f'/api/quotes/{quote["id"]}/history')
        self.assertTrue(any(float(h['snapshot']['items'][0]['quantity']) == 3 for h in history))
        # The source can be ambiguous: preserve null rather than calculating a price.
        corrected['items'][0]['amount'] = None
        missing = self.api('PUT', f'/api/quotes/{quote["id"]}', corrected)
        self.assertIsNone(missing['items'][0]['amount'])
        self.assertEqual(float(missing['items'][0]['computed_amount']), 96000)
        self.assertEqual(self.confirm(missing)['review_status'], 'confirmed')

    def test_06_explicit_contact_title_phone_item_and_period_filters(self):
        quote = self.new_archive()
        quote.update(customer=f'CUSTOMER_{self.token}', contact='김아카이브', phone='02-1234-5678',
                     email='archive@example.invalid', project_name='자동화 설비 보수 프로젝트', issue_date='2026-10-13')
        corrected = self.api('PUT', f'/api/quotes/{quote["id"]}', quote)
        for field, value in [('contact', '김아카이브'), ('phone', '1234-5678'), ('title', '설비 보수')]:
            found = self.api('GET', '/api/quotes?' + urlencode({field: value}))
            self.assertEqual([q['id'] for q in found], [corrected['id']])
            missing = self.api('GET', '/api/quotes?' + urlencode({field: 'MISSING_VALUE_' + self.token}))
            self.assertEqual(missing, [])
        filters = {'customer': quote['customer'], 'contact': '김아카이브', 'phone': '1234', 'title': '설비',
                   'item': '브라켓', 'date_from': '2026-10-01', 'date_to': '2026-10-31', 'min_amount': 40000, 'max_amount': 50000}
        found = self.api('GET', '/api/quotes?' + urlencode(filters))
        self.assertEqual([q['id'] for q in found], [corrected['id']])
        self.confirm(corrected)
        natural = self.api('POST', '/api/natural-search', {'question': '김아카이브 담당자의 견적 보여줘'})
        self.assertTrue(natural['answer'])
        self.assertEqual([q['id'] for q in natural['quotes']], [corrected['id']])
        title = self.api('POST', '/api/natural-search', {'question': '자동화 설비 보수 프로젝트 견적 보여줘'})
        self.assertEqual([q['id'] for q in title['quotes']], [corrected['id']])

    def test_07_compare_includes_pending_documents_components_and_all_options(self):
        imported = self.api('POST', '/api/import-samples')['quotes']
        compared = self.api('GET', '/api/compare')
        for quote in imported:
            rows = [r for r in compared if r['quote_id'] == quote['id']]
            self.assertEqual(len(rows), len(quote['items']))
            self.assertEqual(sum(bool(r['is_component']) for r in rows), sum(bool(i['is_component']) for i in quote['items']))
            self.assertTrue(all('currency' in r and 'specification' in r and 'unit' in r for r in rows))
        xls = next(q for q in imported if q['source_name'].endswith('.xls'))
        options = [r for r in compared if r['quote_id'] == xls['id'] and r.get('option_group')]
        self.assertEqual(len(options), 2)

    def test_08_new_quote_copy_export_template_status_routes_are_disabled(self):
        quote = self.new_archive()
        routes = [('POST', '/api/quotes', {}), ('POST', f'/api/quotes/{quote["id"]}/copy', {}),
                  ('PATCH', f'/api/quotes/{quote["id"]}/status', {'status': 'won'}),
                  ('GET', f'/api/quotes/{quote["id"]}/export?format=xlsx', None),
                  ('GET', f'/api/quotes/{quote["id"]}/export?format=pdf', None),
                  ('GET', '/api/templates', None), ('POST', '/api/templates', {})]
        for method, route, body in routes:
            with self.subTest(route=route, method=method):
                status, _, _ = self.request(method, route, data=body)
                self.assertIn(status, (404, 410))

    def test_09_legacy_generated_rows_are_preserved_but_excluded_from_archive(self):
        original = self.new_archive()
        before = self.api('GET', '/api/stats')
        legacy = copy.deepcopy(original)
        legacy.update(id=str(uuid.uuid4()), number='LEGACY-GENERATED', customer=f'LEGACY_{self.token}', source_url='',
                      source_name='', source_text='', extraction_method='manual', parent_id=None,
                      review_status='confirmed', subtotal=9000000000, total=9900000000, tax=900000000)
        legacy_copy = copy.deepcopy(legacy)
        legacy_copy.update(id=str(uuid.uuid4()), parent_id=original['id'], source_url=original['source_url'], source_name=original['source_name'])
        with closing(sqlite3.connect(self.database)) as connection:
            with connection:
                for quote in (legacy, legacy_copy):
                    connection.execute('INSERT INTO quotes(id,payload,file_hash) VALUES(?,?,?)', (quote['id'], json.dumps(quote, ensure_ascii=False), None))
        found = self.api('GET', '/api/quotes')
        self.assertNotIn(legacy['id'], {q['id'] for q in found})
        self.assertNotIn(legacy_copy['id'], {q['id'] for q in found})
        self.assertEqual(self.api('GET', '/api/quotes?' + urlencode({'query': legacy['customer']})), [])
        natural = self.api('POST', '/api/natural-search', {'question': legacy['customer'] + ' 견적 보여줘'})
        self.assertEqual(natural['quotes'], [])
        self.assertNotIn(legacy['id'], {r['quote_id'] for r in self.api('GET', '/api/compare')})
        after = self.api('GET', '/api/stats')
        self.assertEqual(after['quote_count'], before['quote_count'])
        with closing(sqlite3.connect(self.database)) as connection:
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM quotes WHERE id IN (?,?)', (legacy['id'], legacy_copy['id'])).fetchone()[0], 2)

    def test_10_archive_stats_path_safety_and_restart_persistence(self):
        quote = self.new_archive()
        quote['contact'] = '재시작 담당자'
        quote = self.api('PUT', f'/api/quotes/{quote["id"]}', quote)
        listed = self.api('GET', '/api/quotes')
        stats = self.api('GET', '/api/stats')
        self.assertEqual(stats['quote_count'], len(listed))
        expected_items = sum(len(q['items']) for q in listed)
        self.assertEqual(stats.get('item_count', stats.get('stored_item_count')), expected_items)
        self.assertEqual(stats['pending_count'], sum(q['review_status'] == 'pending' for q in listed))
        self.assertEqual(stats.get('customer_count'), len({q['customer'] for q in listed if q['customer']}))
        for route in ('/api/files/..%2Fapp.py', '/api/files/..%5Capp.py'):
            status, _, _ = self.request('GET', route)
            self.assertIn(status, (400, 403, 404, 422))
        self.server.terminate()
        self.server.wait(timeout=10)
        type(self).server = self.start_server()
        self.wait_ready()
        stored = self.api('GET', f'/api/quotes/{quote["id"]}')
        self.assertEqual(stored['contact'], quote['contact'])
        self.assertEqual(stored['items'], quote['items'])
        self.assertEqual(stored['total'], quote['total'])
        self.assertEqual(stored['source_url'], quote['source_url'])
        after = self.api('GET', '/api/stats')
        self.assertEqual(after['quote_count'], stats['quote_count'])

    def test_11_legacy_source_migration_removes_inferred_tax_but_keeps_corrections(self):
        pdf = next(q for q in self.api('POST', '/api/import-samples')['quotes']
                   if q['source_name'].endswith('.pdf') and q['subtotal'] == 16000000)
        old = copy.deepcopy(pdf)
        old.update(archive_schema_version=1, tax_inferred=True,
                   tax_rate=10, tax=1600000, total=17600000, review_status='confirmed',
                   acknowledged_warnings=True, customer='MIGRATION_HEADER_CORRECTION')
        old['items'][0]['specification'] = '사용자가 정정한 규격'
        old['original_fields'] = {'subtotal': 16000000, 'tax_rate': 10, 'tax': 1600000, 'total': 17600000}
        with closing(sqlite3.connect(self.database)) as connection:
            with connection:
                connection.execute('UPDATE quotes SET payload=? WHERE id=?',
                                   (json.dumps(old, ensure_ascii=False), old['id']))
        self.server.terminate()
        self.server.wait(timeout=10)
        type(self).server = self.start_server()
        self.wait_ready()
        migrated = self.api('GET', f'/api/quotes/{old["id"]}')
        self.assertEqual(migrated['customer'], 'MIGRATION_HEADER_CORRECTION')
        self.assertEqual(migrated['items'][0]['specification'], '사용자가 정정한 규격')
        self.assertEqual(migrated['subtotal'], 16000000)
        for field in ('tax_rate', 'tax', 'total'):
            self.assertIsNone(migrated[field])
            self.assertIsNone(migrated['original_fields'].get(field))
        self.assertEqual(migrated['review_status'], 'pending')
        self.assertFalse(migrated['acknowledged_warnings'])
        history = self.api('GET', f'/api/quotes/{old["id"]}/history')
        self.assertTrue(any(h['snapshot'].get('total') == 17600000 for h in history))
        for entry in history:
            source = entry['snapshot'].get('source_text', '')
            self.assertIsNone(re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', source))
            self.assertIsNone(re.search(r'01[016789][- ]?\d{3,4}[- ]?\d{4}', source))


    def test_12_full_source_content_extra_fields_and_search(self):
        from openpyxl import load_workbook
        from openpyxl.comments import Comment
        book=load_workbook(self.xlsx)
        book.active['B2']='FULL-CONTENT-'+self.token
        sheet=book.create_sheet('추가 조건')
        sheet['A1']='보증기간';sheet['B1']='무상 보증 3년 '+self.token
        sheet['A2']='배송조건';sheet['B2']='지정 장소 반입 '+self.token
        sheet['D7']='=1+2';sheet['D7'].comment=Comment('수식 검토 메모 '+self.token,'가상 담당')
        sheet.row_dimensions[7].hidden=True
        hidden=book.create_sheet('숨김 참고');hidden['C5']='내부 참고 '+self.token;hidden.sheet_state='hidden'
        path=self.work/'full-content.xlsx';book.save(path)
        quote=self.upload([path])['quotes'][0]
        sheets={s['name']:s for s in quote['document_content']['sheets']}
        cell=next(c for c in sheets['추가 조건']['cells'] if c['address']=='D7')
        self.assertEqual(cell['formula'],'=1+2');self.assertIsNone(cell['value'])
        self.assertTrue(cell['row_hidden']);self.assertIn('수식 검토 메모',cell['comment'])
        self.assertEqual(sheets['숨김 참고']['state'],'hidden')
        self.assertTrue(any('무상 보증 3년' in f['value'] for f in quote['extra_fields']))
        self.assertIn(quote['id'],{q['id'] for q in self.api('GET','/api/quotes?'+urlencode({'query':'수식 검토 메모 '+self.token}))})
        original_content=copy.deepcopy(quote['document_content'])
        original_fields=copy.deepcopy(quote['original_extra_fields'])
        quote['extra_fields'].append({'label':'특약','value':'검색 가능한 특약 '+self.token})
        quote['document_content']={'text':'tampered'}
        saved=self.api('PUT',f'/api/quotes/{quote["id"]}',quote)
        self.assertEqual(saved['document_content'],original_content)
        self.assertEqual(saved['original_extra_fields'],original_fields)
        self.assertIn(saved['id'],{q['id'] for q in self.api('GET','/api/quotes?'+urlencode({'query':'검색 가능한 특약 '+self.token}))})
        pdfs=[q for q in self.api('POST','/api/import-samples')['quotes'] if q['source_name'].endswith('.pdf')]
        for pdf in pdfs:
            self.assertTrue(pdf['document_content']['pages'])
            for page in pdf['document_content']['pages']:
                self.assertIn('tables',page);self.assertTrue(page['text'])
        word=self.upload([self.docx])['quotes'][0]
        self.assertTrue(any(b['kind']=='table' for b in word['document_content']['blocks']))

    def test_13_trash_restore_preserves_original_and_excludes_all_active_views(self):
        quote=self.confirm(self.new_archive())
        before_stats=self.api('GET','/api/stats')
        _,source_before,_=self.request('GET',quote['source_url'])
        trashed=self.api('DELETE',f'/api/quotes/{quote["id"]}')
        self.assertTrue(trashed['deleted_at'])
        self.api('GET',f'/api/quotes/{quote["id"]}',expected=404)
        self.assertNotIn(quote['id'],{q['id'] for q in self.api('GET','/api/quotes')})
        self.assertNotIn(quote['id'],{r['quote_id'] for r in self.api('GET','/api/compare')})
        self.assertEqual(self.api('GET','/api/stats')['quote_count'],before_stats['quote_count']-1)
        self.assertIn(quote['id'],{q['id'] for q in self.api('GET','/api/trash')})
        path=self.work/'duplicate-trashed.xlsx';path.write_bytes(source_before)
        result=self.upload([path]);self.assertFalse(result['quotes']);self.assertIn('휴지통',result['errors'][0]['detail'])
        restored=self.api('POST',f'/api/quotes/{quote["id"]}/restore')
        for field in ('items','subtotal','tax','total','extra_fields','document_content','review_status','source_url'):
            self.assertEqual(restored[field],quote[field])
        self.assertFalse(restored.get('deleted_at'))
        self.assertNotIn(quote['id'],{q['id'] for q in self.api('GET','/api/trash')})
        self.assertIn(quote['id'],{q['id'] for q in self.api('GET','/api/quotes')})
        _,source_after,_=self.request('GET',quote['source_url']);self.assertEqual(source_before,source_after)
        history=self.api('GET',f'/api/quotes/{quote["id"]}/history')
        self.assertTrue(any(h['snapshot'].get('deleted_at') for h in history))
        self.assertEqual(self.api('GET','/api/stats')['quote_count'],before_stats['quote_count'])

    def test_14_backfill_preserves_corrections_and_review_status(self):
        quote=self.new_archive();quote['customer']='보강 전 정정 거래처';quote['items'][0]['specification']='정정한 규격'
        quote=self.api('PUT',f'/api/quotes/{quote["id"]}',quote);quote=self.confirm(quote)
        legacy=copy.deepcopy(quote)
        for field in ('document_content','document_content_version','extra_fields','original_extra_fields'):legacy.pop(field,None)
        with closing(sqlite3.connect(self.database)) as connection:
            with connection:connection.execute('UPDATE quotes SET payload=? WHERE id=?',(json.dumps(legacy,ensure_ascii=False),quote['id']))
        self.server.terminate();self.server.wait(timeout=10);type(self).server=self.start_server();self.wait_ready()
        stored=self.api('GET',f'/api/quotes/{quote["id"]}')
        self.assertEqual(stored['document_content_version'],1);self.assertTrue(stored['document_content']['sheets'])
        for field in ('customer','items','total','subtotal','tax','review_status','acknowledged_warnings'):
            self.assertEqual(stored[field],quote[field])
        self.assertTrue(any(h['snapshot'].get('customer')=='보강 전 정정 거래처' for h in self.api('GET',f'/api/quotes/{quote["id"]}/history')))


    def test_15_detailed_facts_are_structured_searchable_and_keep_provenance(self):
        from openpyxl import load_workbook
        book=load_workbook(self.xlsx);book.active['B2']='DETAILED-'+self.token
        sheet=book.create_sheet('회사와 계약조건')
        for row in [('공급사','가상 세부 공급사'),('사업자등록번호','TEST-REG-'+self.token),('보증기간','명시된 보증 24개월'),('배송조건','고객 지정 위치 반입'),('할인율','3%'),('특약','조건 '+self.token)]:sheet.append(row)
        path=self.work/'detailed.xlsx';book.save(path)
        quote=self.upload([path])['quotes'][0]
        warranty=next(f for f in quote['detailed_fields'] if f['key']=='warranty')
        self.assertEqual(warranty['value'],'명시된 보증 24개월')
        self.assertEqual(warranty['source']['sheet'],'회사와 계약조건')
        self.assertEqual(warranty['source']['cell'],'A3');self.assertEqual(warranty['source']['value_cell'],'B3')
        original=copy.deepcopy(quote['original_detailed_fields'])
        warranty['value']='검토 후 수정한 보증 '+self.token
        warranty['source']={'page':999}
        quote['items'][0]['specification']='-. CPU : Intel i7\n-. O/S : Windows 11 Pro\n-. 재질 : SUS304'
        saved=self.api('PUT',f'/api/quotes/{quote["id"]}',quote)
        corrected=next(f for f in saved['detailed_fields'] if f['id']==warranty['id'])
        self.assertEqual(corrected['source']['cell'],'A3');self.assertEqual(corrected['origin'],'corrected')
        self.assertEqual(saved['original_detailed_fields'],original)
        attrs=saved['item_specifications'][0]['attributes']
        self.assertIn(('O/S','Windows 11 Pro'),[(a['label'],a['value']) for a in attrs])
        self.assertIn(saved['id'],{q['id'] for q in self.api('GET','/api/quotes?'+urlencode({'query':'검토 후 수정한 보증 '+self.token}))})
        with closing(sqlite3.connect(self.database)) as connection:
            facts=connection.execute('SELECT category,label,value,item_index,source FROM quote_facts WHERE quote_id=?',(saved['id'],)).fetchall()
        self.assertEqual(len(facts),len(saved['detailed_fields'])+sum(len(it['attributes']) for it in saved['item_specifications']))
        self.assertTrue(any(row[2]=='Windows 11 Pro' and row[3]==0 for row in facts))
        self.assertTrue(any(row[2]=='검토 후 수정한 보증 '+self.token and json.loads(row[4]).get('cell')=='A3' for row in facts))
        bad=copy.deepcopy(saved);bad['detailed_fields'][0]['value']=None
        self.api('PUT',f'/api/quotes/{saved["id"]}',bad,expected=422)


    def test_16_multiple_people_contacts_and_original_provenance(self):
        from openpyxl import load_workbook
        book=load_workbook(self.xlsx);book.active['B2']='PEOPLE-'+self.token
        sheet=book.create_sheet('담당자 명단')
        rows=[('담당자','가상김담당'),('소속회사','가상공급사'),('부서','영업1팀'),('직함','과장'),('전화','02-0000-0000'),('휴대전화','010-0000-0000'),('이메일','sales-'+self.token+'@example.invalid'),('팩스','02-0000-0001'),('주소','가상시 가상로 1'),('참조','가상이담당'),('소속회사','가상고객사'),('부서','구매팀'),('직함','대리'),('이메일','buyer-'+self.token+'@example.invalid')]
        for row in rows:sheet.append(row)
        hidden=book.create_sheet('숨김 인적 사항');hidden.sheet_state='hidden'
        hidden.append(('담당자','가상숨김담당'));hidden.append(('이메일','hidden-person@example.invalid'))
        path=self.work/'people.xlsx';book.save(path)
        quote=self.upload([path])['quotes'][0]
        person=next(p for p in quote['people_records'] if p['name']=='가상김담당')
        for field,value in [('company','가상공급사'),('department','영업1팀'),('title','과장'),('mobile','010-0000-0000'),('fax','02-0000-0001')]:self.assertEqual(person[field],value)
        self.assertTrue(any(p['name']=='가상이담당' and p['department']=='구매팀' for p in quote['people_records']))
        self.assertTrue(any(p['name']=='가상숨김담당' and p['source_records'][0]['source']['hidden'] for p in quote['people_records']))
        original=copy.deepcopy(quote['original_people_records']);evidence=copy.deepcopy(person['source_records'])
        person['email']='corrected-'+self.token+'@example.invalid';person['source_records']=[]
        saved=self.api('PUT',f'/api/quotes/{quote["id"]}',quote)
        corrected=next(p for p in saved['people_records'] if p['id']==person['id'])
        self.assertEqual(corrected['source_records'],evidence);self.assertEqual(saved['original_people_records'],original)
        self.assertIn(saved['id'],{q['id'] for q in self.api('GET','/api/quotes?'+urlencode({'query':person['email']}))})
        self.assertIn(saved['id'],{q['id'] for q in self.api('GET','/api/quotes?'+urlencode({'phone':'01000000000'}))})
        with closing(sqlite3.connect(self.database)) as connection:
            rows=connection.execute('SELECT name,department,email FROM quote_people WHERE quote_id=?',(saved['id'],)).fetchall()
        self.assertIn(('가상김담당','영업1팀',person['email']),rows)

    def test_17_people_pdf_uses_visible_render_not_covered_text(self):
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        pdfmetrics.registerFont(TTFont('PeopleTestKR','C:/Windows/Fonts/malgun.ttf'))
        path=self.work/'people-masked.pdf';doc=canvas.Canvas(str(path),pagesize=(700,700));doc.setFont('PeopleTestKR',18)
        doc.drawString(35,650,'견적번호 : MASK-'+self.token)
        doc.drawString(35,610,'담당자 : 가상공개담당')
        doc.drawString(35,570,'이메일 : visible-'+self.token+'@example.invalid')
        doc.drawString(35,530,'전화 : 02-0000-0000')
        hidden='hidden-'+self.token+'@example.invalid'
        doc.drawString(35,450,'이메일 : '+hidden)
        doc.setFillColorRGB(0,0,0);doc.rect(25,425,640,55,fill=1,stroke=0)
        doc.showPage();doc.save()
        quote=self.upload([path])['quotes'][0]
        people_json=json.dumps(quote['people_records'],ensure_ascii=False)
        evidence_json=json.dumps(quote['personal_source_blocks'],ensure_ascii=False)
        self.assertNotIn(hidden,people_json);self.assertNotIn(hidden,evidence_json)
        self.assertTrue(quote['personal_source_blocks'])
        self.assertTrue(all(block['source'].get('method')=='visible-ocr' for block in quote['personal_source_blocks']))
        self.assertTrue(any('가상공개담당' in p['name'].replace(' ','') for p in quote['people_records']),people_json+'\n'+json.dumps(quote.get('personal_visible_pages'),ensure_ascii=False))


    def test_18_multiple_names_on_same_line_and_unassigned_contacts(self):
        import importlib.util
        spec=importlib.util.spec_from_file_location('people_capture_unit',ROOT/'backend'/'people_capture.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        people,evidence=module.parse_blocks([{'text':'전화:02-0000-0000\n담당자:가상김담당 이메일:supplier@example.invalid / 참조:가상이담당 이메일:buyer@example.invalid','source':{'page':1,'method':'visible-ocr'}}])
        self.assertTrue(any(not p['name'] and p['phone']=='02-0000-0000' for p in people))
        self.assertTrue(any(p['name']=='가상김담당' and 'supplier@example.invalid' in p['email'] for p in people))
        self.assertTrue(any(p['name']=='가상이담당' and 'buyer@example.invalid' in p['email'] for p in people))
        self.assertTrue(all(p['review_status']=='pending' for p in people))
        rows,_=module.parse_blocks([{'text':'참조 허웅프로님 귀하 견적일자\n견적담당자 연 락 처 E - M AIL [PAYMENT]\n가상김이사 010-0000-0000','source':{'page':1}}])
        self.assertTrue(any(p['name']=='허웅' and p['title']=='프로' for p in rows))
        self.assertTrue(any(p['name']=='가상김' and p['title']=='이사' for p in rows))
        self.assertFalse(any('연 락 처' in p['name'] for p in rows))


    def test_19_category_lists_values_owners_search_and_trash(self):
        quote=self.new_archive()
        people=[{'id':uuid.uuid4().hex,'name':'가상카테고리담당','company':'가상카테고리회사','department':'구매팀','title':'과장','role':'고객 참조','phone':'02-0000-0011','mobile':'010-0000-0011','fax':'02-0000-0022','email':'category-'+self.token+'@example.invalid','address':'가상 주소'}]
        quote['people_records']=people
        quote['extra_fields'].append({'label':'FAX','value':'02-0000-0022 M : 010-0000-0011 MAIL : other@example.invalid'})
        quote=self.api('PUT',f'/api/quotes/{quote["id"]}',quote)
        catalog=self.api('GET','/api/categories')
        self.assertIn('fax',{c['key'] for c in catalog})
        result=self.api('GET','/api/categories/fax?'+urlencode({'query':'가상카테고리담당'}))
        entries=[row for row in result['values'] if row['quote_id']==quote['id']]
        named=[row for row in entries if row['person_name']=='가상카테고리담당']
        self.assertEqual(len(named),1);self.assertEqual(named[0]['value'],'02-0000-0022')
        self.assertTrue(all('MAIL' not in row['value'] and 'M :' not in row['value'] for row in entries))
        self.assertEqual(entries[0]['person_name'],'가상카테고리담당')
        numeric=self.api('GET','/api/categories/fax?query=0200000022')
        self.assertIn(quote['id'],{r['quote_id'] for r in numeric['values']})
        self.assertIn(quote['id'],{r['quote_id'] for r in self.api('GET','/api/categories/email')['values']})
        self.api('GET','/api/categories/not-a-category',expected=404)
        self.api('DELETE',f'/api/quotes/{quote["id"]}')
        self.assertNotIn(quote['id'],{r['quote_id'] for r in self.api('GET','/api/categories/fax')['values']})
        self.api('POST',f'/api/quotes/{quote["id"]}/restore')
        self.assertIn(quote['id'],{r['quote_id'] for r in self.api('GET','/api/categories/fax')['values']})


    def test_20_api_credentials_require_user_consent_and_never_enter_database(self):
        body,content_type=self.multipart([(self.xlsx.name,self.xlsx.read_bytes())])
        secret='FAKE_REQUEST_ONLY_TOKEN_12345'
        request=Request(self.base+'/api/upload',data=body,method='POST',headers={'Content-Type':content_type,'X-User-API-Key':secret,'X-User-Agent-ID':'agt_mock0000'})
        with self.assertRaises(HTTPError) as error:urlopen(request,timeout=10)
        self.assertEqual(error.exception.code,422)
        response=error.exception.read().decode();self.assertNotIn(secret,response)
        with closing(sqlite3.connect(self.database)) as connection:
            self.assertEqual(connection.execute('SELECT count(*) FROM quotes WHERE payload LIKE ?',('%'+secret+'%',)).fetchone()[0],0)
            self.assertEqual(connection.execute('SELECT count(*) FROM quote_history WHERE snapshot LIKE ?',('%'+secret+'%',)).fetchone()[0],0)
        self.assertFalse(self.api('GET','/api/health')['shared_api_key'])


class RecordingResult(unittest.TextTestResult):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.records = []

    def startTest(self, test):
        self.started = time.monotonic()
        super().startTest(test)

    def addSuccess(self, test):
        self.records.append({'test': test.id(), 'status': 'passed', 'seconds': round(time.monotonic() - self.started, 3)})
        super().addSuccess(test)

    def addFailure(self, test, err):
        self.records.append({'test': test.id(), 'status': 'failed', 'detail': self._exc_info_to_string(err, test)})
        super().addFailure(test, err)

    def addError(self, test, err):
        self.records.append({'test': str(test), 'status': 'error', 'detail': self._exc_info_to_string(err, test)})
        super().addError(test, err)

    def addSubTest(self, test, subtest, err):
        if err is not None:
            self.records.append({'test': str(subtest), 'status': 'failed', 'detail': self._exc_info_to_string(err, test)})
        super().addSubTest(test, subtest, err)


if __name__ == '__main__':
    started = time.monotonic()
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(QuoteIntegration)
    result = unittest.TextTestRunner(verbosity=2, resultclass=RecordingResult).run(suite)
    report = {'scope': 'Existing quote archive: import, preserve original values, correct extraction, search and view original',
              'ok': result.wasSuccessful(), 'tests_run': result.testsRun, 'failures': len(result.failures), 'errors': len(result.errors),
              'duration_seconds': round(time.monotonic() - started, 2),
              'isolation': 'Temporary QUOTE_DATA_DIR and disposable SQLite; external AI disabled; originals read-only', 'tests': result.records}
    (ROOT / 'tests' / 'integration-results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    sys.exit(0 if result.wasSuccessful() else 1)
