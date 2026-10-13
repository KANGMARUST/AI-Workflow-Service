"""Real HTTP integration checks with a disposable database and file store.

Run with the bundled Python: python tests/test_integration.py
The real quote files are read from project/data; they are never modified.
"""
from __future__ import annotations

import copy
import io
import json
import os
from pathlib import Path
import re
import socket
import subprocess
import sys
import tempfile
import time
import unittest
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
import uuid

ROOT = Path(__file__).resolve().parents[1]
PACKAGES = ROOT / 'backend' / '.packages'
if PACKAGES.exists():
    sys.path.insert(0, str(PACKAGES))


class QuoteIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.run_root = (ROOT / 'tests' / '.runs').resolve()
        cls.run_root.mkdir(exist_ok=True)
        cls.temporary = tempfile.TemporaryDirectory(prefix='quote-integration-', dir=cls.run_root)
        cls.work = Path(cls.temporary.name)
        cls.token = uuid.uuid4().hex[:8]
        cls.counter = 0
        with socket.socket() as listener:
            listener.bind(('127.0.0.1', 0))
            cls.port = listener.getsockname()[1]
        cls.base = f'http://127.0.0.1:{cls.port}'
        env = os.environ.copy()
        env['QUOTE_DATA_DIR'] = str(cls.work / 'storage')
        env['QUOTE_SAMPLE_DIR'] = str(ROOT / 'data')
        env['EXTERNAL_AI_ENABLED'] = 'false'
        # These tests must never connect to an existing database or remote AI.
        env['DATABASE_URL'] = 'sqlite:///' + (cls.work / 'storage' / 'quotes.db').as_posix()
        env.pop('OPENAI_API_KEY', None)
        env['PYTHONPATH'] = os.pathsep.join([str(PACKAGES), str(ROOT / 'backend'), env.get('PYTHONPATH', '')])
        cls.environment = env
        cls.log = open(cls.work / 'server.log', 'w', encoding='utf-8')
        cls.server = cls.start_server()
        cls.addClassCleanup(cls.cleanup)
        cls.wait_ready()
        cls.make_fixtures()

    @classmethod
    def start_server(cls):
        return subprocess.Popen(
            [sys.executable, '-m', 'uvicorn', 'app:app', '--host', '127.0.0.1', '--port', str(cls.port)],
            cwd=ROOT / 'backend', env=cls.environment, stdout=cls.log, stderr=subprocess.STDOUT,
            creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0),
        )

    @classmethod
    def wait_ready(cls):
        deadline = time.monotonic() + 45
        while time.monotonic() < deadline:
            try:
                status, health, _ = cls.request('GET', '/api/health')
                if status == 200 and health.get('ok'):
                    return
            except (OSError, ValueError):
                pass
            if cls.server.poll() is not None:
                break
            time.sleep(0.2)
        cls.log.flush()
        raise RuntimeError('Isolated API did not become ready: ' + (cls.work / 'server.log').read_text(encoding='utf-8')[-6000:])

    @classmethod
    def cleanup(cls):
        if getattr(cls, 'server', None) and cls.server.poll() is None:
            cls.server.terminate()
            try:
                cls.server.wait(timeout=10)
            except subprocess.TimeoutExpired:
                cls.server.kill()
                cls.server.wait(timeout=5)
        if getattr(cls, 'log', None):
            cls.log.close()
        if getattr(cls, 'temporary', None):
            if not cls.work.resolve().is_relative_to(cls.run_root):
                raise RuntimeError('Refusing to clean a temporary directory outside tests/.runs')
            cls.temporary.cleanup()

    @classmethod
    def request(cls, method, route, data=None, body=None, content_type=None):
        headers = {}
        if data is not None:
            body = json.dumps(data, ensure_ascii=False).encode('utf-8')
            headers['Content-Type'] = 'application/json'
        elif content_type:
            headers['Content-Type'] = content_type
        req = Request(cls.base + route, data=body, method=method, headers=headers)
        try:
            response = urlopen(req, timeout=60)
        except HTTPError as exc:
            response = exc
        with response:
            raw = response.read()
            content = response.headers.get('Content-Type', '')
            parsed = json.loads(raw) if 'json' in content and raw else raw
            return response.status, parsed, {key.lower(): value for key, value in response.headers.items()}

    def api(self, method, route, data=None, expected=200):
        status, result, _ = self.request(method, route, data=data)
        self.assertEqual(status, expected, f'{method} {route}: {result}')
        return result

    @classmethod
    def multipart(cls, files, fields=None):
        boundary = 'QuoteTest' + uuid.uuid4().hex
        parts = []
        for key, value in (fields or {}).items():
            parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{key}"\r\n\r\n{value}\r\n'.encode('utf-8'))
        for field, filename, mime, data in files:
            parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{field}"; filename="{filename}"\r\nContent-Type: {mime}\r\n\r\n'.encode('utf-8') + data + b'\r\n')
        parts.append(f'--{boundary}--\r\n'.encode('ascii'))
        return b''.join(parts), f'multipart/form-data; boundary={boundary}'

    def upload(self, paths, route='/api/upload', field='files', fields=None):
        mime = {'.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', '.png': 'image/png', '.txt': 'text/plain'}
        body, content = self.multipart([(field, p.name, mime.get(p.suffix, 'application/octet-stream'), p.read_bytes()) for p in paths], fields)
        status, result, _ = self.request('POST', route, body=body, content_type=content)
        self.assertEqual(status, 200, str(result))
        return result

    @classmethod
    def make_fixtures(cls):
        from openpyxl import Workbook
        from docx import Document
        from PIL import Image, ImageDraw, ImageFont
        rows = [
            ['견적서'], ['견적번호', 'FIXTURE-' + cls.token], ['거래처', '테스트 고객'],
            ['작성일', '2026-10-13'], ['통화', 'KRW'], ['부가세율', 10],
            ['품목명', '규격', '단위', '수량', '단가', '금액'],
            ['브라켓', '100mm', 'EA', 3, 12000, 36000],
            ['볼트', 'M8', 'EA', 5, 1000, 5000],
            ['공급가액', 41000], ['부가세', 4100], ['총액', 45100],
        ]
        book = Workbook()
        for row in rows:
            book.active.append(row)
        cls.xlsx = cls.work / 'fixture.xlsx'
        book.save(cls.xlsx)
        doc = Document()
        for row in rows[:6]:
            doc.add_paragraph(' : '.join(map(str, row)))
        table = doc.add_table(rows=1, cols=6)
        for cell, label in zip(table.rows[0].cells, rows[6]):
            cell.text = label
        for row in rows[7:9]:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = str(value)
        for row in rows[9:]:
            doc.add_paragraph(' : '.join(map(str, row)))
        cls.docx = cls.work / 'fixture.docx'
        doc.save(cls.docx)
        picture = Image.new('RGB', (1800, 1000), 'white')
        draw = ImageDraw.Draw(picture)
        font_path = Path('C:/Windows/Fonts/malgun.ttf')
        font = ImageFont.truetype(str(font_path), 38) if font_path.exists() else ImageFont.load_default()
        image_lines = [
            '견적서', '견적번호 : OCR-' + cls.token, '거래처 : 테스트 고객',
            '작성일 : 2026-10-13', '품목명    규격    단위    수량    단가    금액',
            '브라켓    100mm    EA    3    12000    36000',
            '공급가액 : 36000', '부가세 : 3600', '총액 : 39600',
        ]
        for i, line in enumerate(image_lines):
            draw.text((60, 35 + i * 90), line, fill='black', font=font)
        cls.png = cls.work / 'fixture.png'
        picture.save(cls.png)
        cls.invalid = cls.work / 'unsupported.txt'
        cls.invalid.write_text('This unsupported file must produce a per-file error.', encoding='utf-8')

    def draft(self, **changes):
        type(self).counter += 1
        draft = {
            'number': f'TEST-{self.token}-{self.counter}', 'customer': f'TEST_CUSTOMER_{self.token}',
            'issuer': '테스트 발행사', 'issue_date': '2026-10-13', 'currency': 'KRW',
            'tax_rate': 10, 'discount': 0, 'delivery': '2026-11-01',
            'payment_terms': '검수 후 30일', 'items': [
                {'name': 'TEST_BRACKET', 'specification': '100mm', 'unit': 'EA',
                 'quantity': 3, 'unit_price': 12000, 'amount': 36000,
                 'is_component': False, 'parent_index': None, 'original_text': ''},
            ],
        }
        draft.update(changes)
        return self.api('POST', '/api/quotes', draft)

    def confirm(self, quote):
        return self.api('POST', f'/api/quotes/{quote["id"]}/confirm', {'acknowledged_warnings': True})

    def test_01_real_samples_and_duplicate_import(self):
        imported = self.api('POST', '/api/import-samples')
        self.assertEqual(imported['errors'], [])
        self.assertEqual(len(imported['quotes']), 3)
        pdfs = [q for q in imported['quotes'] if q['source_name'].lower().endswith('.pdf')]
        self.assertEqual(len(pdfs), 2)
        by_amount = {round(float(q['subtotal'])): q for q in pdfs}
        self.assertIn(16000000, by_amount)
        self.assertIn(266740000, by_amount)
        first = by_amount[16000000]
        self.assertEqual(len(first['items']), 3)
        complex_quote = by_amount[266740000]
        self.assertEqual(sum(not item['is_component'] for item in complex_quote['items']), 5)
        components = [item for item in complex_quote['items'] if item['is_component']]
        self.assertGreater(len(components), 0)
        for item in components:
            self.assertIsInstance(item['parent_index'], int)
            parent = complex_quote['items'][item['parent_index']]
            self.assertFalse(parent['is_component'])
        second = self.api('POST', '/api/import-samples')
        self.assertEqual({q['id'] for q in imported['quotes']}, {q['id'] for q in second['quotes']})
        for quote in imported['quotes']:
            self.assertEqual(quote['review_status'], 'pending')
            self.assertTrue(quote['source_url'])
            status, file_content, _ = self.request('GET', quote['source_url'])
            self.assertEqual(status, 200)
            self.assertGreater(len(file_content), 0)
            # Redacted source documents must not leak hidden PDF contact text.
            self.assertIsNone(re.search(r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}', quote['source_text']))
            self.assertIsNone(re.search(r'01[016789][- ]?\d{3,4}[- ]?\d{4}', quote['source_text']))
        xls = next(q for q in imported['quotes'] if q['source_name'].lower().endswith('.xls'))
        self.assertEqual(float(xls['subtotal']), 6750000)
        options = [item for item in xls['items'] if item.get('option_group')]
        self.assertEqual(len(options), 2)
        self.assertFalse(any(item.get('included', True) for item in options))
        self.assertEqual(len({item['option_group'] for item in options}), 1)
        status, _, _ = self.request('POST', f'/api/quotes/{xls["id"]}/confirm', data={'acknowledged_warnings': True})
        self.assertIn(status, (400, 409, 422))
        # Supply required headers without overwriting source data files.
        xls['customer'] = xls['customer'] or '선택옵션 테스트 고객'
        xls['issue_date'] = xls['issue_date'] or '2026-10-13'
        options[0]['included'] = True
        changed = self.api('PUT', f'/api/quotes/{xls["id"]}', xls)
        self.assertIn(float(changed['subtotal']), (9660000, 9070000))
        self.assertEqual(self.confirm(changed)['review_status'], 'confirmed')
        for item in changed['items']:
            if item.get('option_group'):
                item['included'] = True
        both = self.api('PUT', f'/api/quotes/{xls["id"]}', changed)
        status, _, _ = self.request('POST', f'/api/quotes/{both["id"]}/confirm', data={'acknowledged_warnings': True})
        self.assertIn(status, (400, 409, 422))

    def test_02_xlsx_docx_upload_and_partial_error(self):
        result = self.upload([self.xlsx, self.docx, self.invalid])
        self.assertEqual(len(result['quotes']), 2)
        self.assertEqual(len(result['errors']), 1)
        for quote in result['quotes']:
            self.assertEqual(quote['review_status'], 'pending')
            self.assertEqual(len(quote['items']), 2, quote)
            self.assertEqual(float(quote['subtotal']), 41000)
            self.assertEqual(float(quote['total']), 45100)
            self.assertEqual(quote['customer'], '테스트 고객')
        duplicate = self.upload([self.xlsx])
        self.assertEqual(duplicate['quotes'][0]['id'], result['quotes'][0]['id'])

    def test_03_real_ocr_image(self):
        health = self.api('GET', '/api/health')
        self.assertTrue(health['ocr_available'], 'Prototype requires an installed, functioning local OCR engine')
        result = self.upload([self.png])
        self.assertEqual(result['errors'], [])
        self.assertEqual(len(result['quotes']), 1)
        quote = result['quotes'][0]
        self.assertIn('ocr', quote['extraction_method'].lower())
        self.assertIn('36000', quote['source_text'].replace(',', '').replace(' ', ''))
        self.assertGreater(len(quote['items']), 0, quote['source_text'])
        bracket = next((item for item in quote['items'] if '브라켓' in item['name']), None)
        self.assertIsNotNone(bracket, quote['source_text'])
        self.assertIn('100mm', bracket['specification'])
        self.assertEqual(float(bracket['quantity']), 3)
        self.assertEqual(float(bracket['unit_price']), 12000)
        self.assertEqual(float(bracket['amount']), 36000)

    def test_04_review_copy_history_and_exports(self):
        draft = self.draft()
        status, _, _ = self.request('GET', f'/api/quotes/{draft["id"]}/export?format=xlsx')
        self.assertIn(status, (400, 409, 422))
        confirmed = self.confirm(draft)
        self.assertEqual(confirmed['review_status'], 'confirmed')
        duplicate = self.api('POST', f'/api/quotes/{draft["id"]}/copy')
        self.assertNotEqual(duplicate['id'], draft['id'])
        self.assertNotEqual(duplicate['number'], draft['number'])
        self.assertEqual(duplicate['parent_id'], draft['id'])
        self.assertEqual(duplicate['review_status'], 'pending')
        before = copy.deepcopy(duplicate)
        duplicate['items'][0]['quantity'] = 8
        duplicate['items'][0]['amount'] = 96000
        changed = self.api('PUT', f'/api/quotes/{duplicate["id"]}', duplicate)
        self.assertEqual(float(changed['subtotal']), 96000)
        self.assertEqual(float(changed['total']), 105600)
        self.assertEqual(changed['review_status'], 'pending')
        self.assertGreater(changed['version'], before['version'])
        history = self.api('GET', f'/api/quotes/{duplicate["id"]}/history')
        self.assertTrue(any(float(entry['snapshot']['items'][0]['quantity']) == 3 for entry in history))
        self.confirm(changed)
        from openpyxl import load_workbook
        for template in ('standard', 'compact'):
            status, content, headers = self.request('GET', f'/api/quotes/{duplicate["id"]}/export?' + urlencode({'format': 'xlsx', 'template': template}))
            self.assertEqual(status, 200)
            book = load_workbook(io.BytesIO(content), data_only=True)
            values = [cell for sheet in book for row in sheet.iter_rows(values_only=True) for cell in row]
            self.assertIn('TEST_BRACKET', values)
            self.assertIn(8, values)
            self.assertIn(105600, values)
            self.assertIn('spreadsheet', headers.get('content-type', ''))
        import pdfplumber
        pdf_texts = {}
        for template in ('standard', 'compact'):
            status, content, _ = self.request('GET', f'/api/quotes/{duplicate["id"]}/export?' + urlencode({'format': 'pdf', 'template': template}))
            self.assertEqual(status, 200)
            self.assertTrue(content.startswith(b'%PDF'))
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text = ''.join(page.extract_text() or '' for page in pdf.pages)
            self.assertIn('TEST_BRACKET', text)
            self.assertIn('105600', text.replace(',', '').replace(' ', ''))
            pdf_texts[template] = text
        self.assertIn('견 적 서', pdf_texts['standard'])
        self.assertIn('견적 요약서', pdf_texts['compact'])
        self.assertNotEqual(pdf_texts['standard'], pdf_texts['compact'])
        # Modifying a confirmed quote invalidates review confirmation.
        latest = self.api('GET', f'/api/quotes/{duplicate["id"]}')
        latest['notes'] = '검토 후 수정 테스트'
        edited = self.api('PUT', f'/api/quotes/{duplicate["id"]}', latest)
        self.assertEqual(edited['review_status'], 'pending')

    def test_05_blocking_validation_and_warning_acknowledgement(self):
        invalid_changes = [
            {'customer': ''}, {'issue_date': ''}, {'items': []},
            {'items': [{'name': 'A', 'quantity': None, 'unit_price': 10, 'amount': None}]},
            {'items': [{'name': 'A', 'quantity': -1, 'unit_price': 10, 'amount': -10}]},
            {'items': [{'name': 'A', 'quantity': 1, 'unit_price': -10, 'amount': -10}]},
            {'items': [{'name': 'A', 'quantity': 2, 'unit_price': 10, 'amount': 999}]},
            {'discount': 100000}, {'tax_rate': -1}, {'tax_rate': 101},
        ]
        for changes in invalid_changes:
            with self.subTest(changes=changes):
                quote = self.draft(**changes)
                self.assertTrue(quote['warnings'])
                status, _, _ = self.request('POST', f'/api/quotes/{quote["id"]}/confirm', data={'acknowledged_warnings': True})
                self.assertIn(status, (400, 409, 422))
        mismatch = self.draft(source_total=1)
        status, _, _ = self.request('POST', f'/api/quotes/{mismatch["id"]}/confirm', data={})
        self.assertIn(status, (400, 409, 422))
        self.assertEqual(self.confirm(mismatch)['review_status'], 'confirmed')

    def test_06_decimal_discount_tax_and_package_totals(self):
        quote = self.draft(currency='USD', tax_rate=10, discount=0.11, items=[
            {'name': 'Package', 'specification': 'A', 'unit': 'SET', 'quantity': 3, 'unit_price': 0.1, 'amount': 0.3, 'is_component': False, 'parent_index': None},
            {'name': 'Included component', 'specification': 'B', 'unit': 'EA', 'quantity': 1, 'unit_price': 999, 'amount': 999, 'is_component': True, 'parent_index': 0},
            {'name': 'Free line', 'quantity': 0, 'unit_price': 50, 'amount': 0, 'is_component': False, 'parent_index': None},
        ])
        self.assertAlmostEqual(float(quote['subtotal']), 0.30, places=2)
        self.assertAlmostEqual(float(quote['tax']), 0.02, places=2)
        self.assertAlmostEqual(float(quote['total']), 0.21, places=2)
        self.assertEqual(self.confirm(quote)['review_status'], 'confirmed')

    def test_07_filters_compare_status_stats_natural_search(self):
        customer = f'NATURAL_{self.token}'
        first = self.confirm(self.draft(customer=customer, issue_date='2025-01-15'))
        second = self.confirm(self.draft(customer=customer, currency='USD', issue_date='2026-10-13', items=[
            {'name': 'TEST_BRACKET', 'specification': '200mm', 'unit': 'EA', 'quantity': 3, 'unit_price': 30, 'amount': 90, 'is_component': False, 'parent_index': None},
        ]))
        pending = self.draft(customer=customer)
        found = self.api('GET', '/api/quotes?' + urlencode({'customer': customer, 'item': 'TEST_BRACKET', 'date_from': '2025-01-01', 'date_to': '2025-12-31', 'min_amount': 30000, 'max_amount': 50000, 'review_status': 'confirmed'}))
        self.assertEqual([q['id'] for q in found], [first['id']])
        compare = self.api('GET', '/api/compare?' + urlencode({'query': 'TEST_BRACKET'}))
        selected = [row for row in compare if row['customer'] == customer]
        self.assertEqual({row['quote_id'] for row in selected}, {first['id'], second['id']})
        self.assertEqual({row['currency'] for row in selected}, {'USD', 'KRW'})
        self.assertEqual({row['specification'] for row in selected}, {'100mm', '200mm'})
        self.assertNotIn(pending['id'], {row['quote_id'] for row in compare})
        won = self.api('PATCH', f'/api/quotes/{first["id"]}/status', {'status': 'won'})
        self.assertEqual(won['status'], 'won')
        filtered = self.api('GET', '/api/quotes?' + urlencode({'customer': customer, 'status': 'won'}))
        self.assertEqual([q['id'] for q in filtered], [first['id']])
        status, _, _ = self.request('PATCH', f'/api/quotes/{first["id"]}/status', data={'status': 'invented'})
        self.assertIn(status, (400, 422))
        stats = self.api('GET', '/api/stats')
        self.assertGreaterEqual(stats['confirmed_count'], 2)
        self.assertGreaterEqual(stats['status_counts'].get('won', 0), 1)
        self.assertIsInstance(stats['monthly'], list)
        self.assertIsInstance(stats['customers'], list)
        natural = self.api('POST', '/api/natural-search', {'question': f'{customer} 견적 보여줘'})
        self.assertTrue(natural['answer'])
        self.assertIsInstance(natural['filters'], dict)
        self.assertTrue(any(q['id'] == first['id'] for q in natural['quotes']))
        self.assertTrue(all(q['customer'] == customer for q in natural['quotes']))

    def test_08_custom_xlsx_template(self):
        from openpyxl import Workbook, load_workbook
        book = Workbook()
        book.active.title = '견적서'
        book.active.append(['거래처', '{{customer}}'])
        book.active.append(['견적번호', '{{number}}'])
        book.active.append(['총액', '{{total}}'])
        book.active.append(['CUSTOM_QUOTE_TITLE', '{{project_name}}'])
        book.create_sheet('items').append(['품목명', '규격', '수량', '단가', '금액'])
        path = self.work / 'custom-template.xlsx'
        book.save(path)
        template = self.upload([path], '/api/templates', 'file', {'name': '테스트 양식'})
        self.assertTrue(template['id'])
        templates = self.api('GET', '/api/templates')
        self.assertTrue({'standard', 'compact', template['id']} <= {t['id'] for t in templates})
        quote = self.confirm(self.draft(project_name='사용자양식 전용 프로젝트'))
        status, content, _ = self.request('GET', f'/api/quotes/{quote["id"]}/export?' + urlencode({'format': 'xlsx', 'template': template['id']}))
        self.assertEqual(status, 200)
        exported = load_workbook(io.BytesIO(content), data_only=True)
        self.assertEqual(exported['견적서']['B1'].value, quote['customer'])
        self.assertEqual(exported['견적서']['B2'].value, quote['number'])
        self.assertEqual(float(exported['견적서']['B3'].value), float(quote['total']))
        self.assertIn('TEST_BRACKET', [cell for row in exported['items'].iter_rows(values_only=True) for cell in row])
        import pdfplumber
        status, pdf_bytes, _ = self.request('GET', f'/api/quotes/{quote["id"]}/export?' + urlencode({'format': 'pdf', 'template': template['id']}))
        self.assertEqual(status, 200)
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as document:
            pdf_text = '\n'.join(page.extract_text() or '' for page in document.pages)
        self.assertIn(quote['customer'], pdf_text)
        self.assertIn('TEST_BRACKET', pdf_text)
        self.assertIn(str(int(quote['total'])), pdf_text.replace(',', ''))
        self.assertIn('테스트 양식', pdf_text)
        self.assertIn('CUSTOM_QUOTE_TITLE', pdf_text)
        self.assertIn('사용자양식 전용 프로젝트', pdf_text)
        self.assertNotIn('{{', pdf_text)

    def test_09_path_traversal_and_restart_persistence(self):
        for route in ('/api/files/..%2Fapp.py', '/api/files/..%5Capp.py'):
            status, _, _ = self.request('GET', route)
            self.assertIn(status, (400, 403, 404, 422))
        quote = self.draft(customer=f'PERSIST_{self.token}')
        self.server.terminate()
        self.server.wait(timeout=10)
        type(self).server = self.start_server()
        self.wait_ready()
        stored = self.api('GET', f'/api/quotes/{quote["id"]}')
        self.assertEqual(stored['customer'], quote['customer'])
        self.assertEqual(stored['items'], quote['items'])

    def test_10_multi_page_pdf_includes_all_items_and_conditions(self):
        import pdfplumber
        import pypdfium2
        items = [
            {'name': f'LONG_ITEM_{index:03d}', 'specification': f'규격 {index} / 100mm',
             'unit': 'EA', 'quantity': index + 1, 'unit_price': 1000, 'amount': (index + 1) * 1000,
             'is_component': False, 'parent_index': None}
            for index in range(80)
        ]
        quote = self.confirm(self.draft(items=items))
        status, content, _ = self.request('GET', f'/api/quotes/{quote["id"]}/export?format=pdf')
        self.assertEqual(status, 200)
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            self.assertGreater(len(pdf.pages), 1)
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        for item in items:
            self.assertIn(item['name'], text)
        self.assertIn('검수 후 30일', text)
        self.assertIn('2026-11-01', text)
        self.assertIn('규격', text)
        # Synthetic quote only: save first/last pages for the main agent's visual QA.
        output = ROOT / 'tests' / 'qa'
        output.mkdir(exist_ok=True)
        document = pypdfium2.PdfDocument(content)
        try:
            for page_index, label in ((0, 'first'), (len(document) - 1, 'last')):
                page = document[page_index]
                try:
                    image = page.render(scale=1.3).to_pil()
                    image.save(output / f'multipage-{label}.png')
                finally:
                    page.close()
        finally:
            document.close()

    def test_11_foreign_currency_stats_are_not_added_to_krw(self):
        before = self.api('GET', '/api/stats')
        foreign = self.confirm(self.draft(currency='USD', customer=f'FOREIGN_ONLY_{self.token}', items=[
            {'name': 'FOREIGN_HIGH_AMOUNT', 'specification': 'USD only', 'unit': 'EA',
             'quantity': 1, 'unit_price': 1000000000, 'amount': 1000000000,
             'is_component': False, 'parent_index': None}
        ]))
        after = self.api('GET', '/api/stats')
        self.assertEqual(after['currency'], 'KRW')
        self.assertEqual(after['quoted_amount'], before['quoted_amount'])
        self.assertEqual(after['foreign_currency_count'], before['foreign_currency_count'] + 1)
        self.assertEqual(after['currency_totals']['USD'], before['currency_totals'].get('USD', 0) + foreign['total'])
        self.assertEqual(sum(month['amount'] for month in after['monthly']), sum(month['amount'] for month in before['monthly']))
        customer = next(c for c in after['customers'] if c['customer'] == foreign['customer'])
        self.assertEqual(customer['amount'], 0)
        self.assertEqual(customer['count'], 1)

    def test_12_forged_review_status_cannot_bypass_confirm_or_export(self):
        forged = self.draft(review_status='confirmed', acknowledged_warnings=True)
        self.assertEqual(forged['review_status'], 'pending')
        status, _, _ = self.request('GET', f'/api/quotes/{forged["id"]}/export?format=pdf')
        self.assertIn(status, (400, 409, 422))
        legitimate = self.confirm(forged)
        legitimate['review_status'] = 'confirmed'
        legitimate['items'][0]['quantity'] = -8
        legitimate['items'][0]['amount'] = -96000
        legitimate['acknowledged_warnings'] = True
        manipulated = self.api('PUT', f'/api/quotes/{legitimate["id"]}', legitimate)
        self.assertEqual(manipulated['review_status'], 'pending')
        self.assertTrue(manipulated['blocking_warnings'])
        status, _, _ = self.request('POST', f'/api/quotes/{manipulated["id"]}/confirm', data={'acknowledged_warnings': True})
        self.assertIn(status, (400, 409, 422))
        for extension in ('xlsx', 'pdf'):
            status, _, _ = self.request('GET', f'/api/quotes/{manipulated["id"]}/export?format={extension}')
            self.assertIn(status, (400, 409, 422))


class RecordingResult(unittest.TextTestResult):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.records = []

    def startTest(self, test):
        self.started = time.monotonic()
        super().startTest(test)

    def addSuccess(self, test):
        self.records.append({'test': test.id(), 'status': 'passed', 'seconds': round(time.monotonic() - self.started, 3)})
        super().addSuccess(test)

    def addFailure(self, test, err):
        self.records.append({'test': test.id(), 'status': 'failed', 'detail': self._exc_info_to_string(err, test)})
        super().addFailure(test, err)

    def addError(self, test, err):
        self.records.append({'test': str(test), 'status': 'error', 'detail': self._exc_info_to_string(err, test)})
        super().addError(test, err)


if __name__ == '__main__':
    started = time.monotonic()
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(QuoteIntegration)
    result = unittest.TextTestRunner(verbosity=2, resultclass=RecordingResult).run(suite)
    report = {'ok': result.wasSuccessful(), 'tests_run': result.testsRun, 'failures': len(result.failures), 'errors': len(result.errors), 'duration_seconds': round(time.monotonic() - started, 2), 'isolation': 'Temporary QUOTE_DATA_DIR and explicit disposable SQLite DATABASE_URL; EXTERNAL_AI_ENABLED=false', 'tests': result.records}
    (ROOT / 'tests' / 'integration-results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    sys.exit(0 if result.wasSuccessful() else 1)
